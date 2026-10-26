"""
毕业论文深度学习内容补写脚本

在已完善的论文（_完善.docx）基础上，补写深度学习相关章节内容：
  1) 第2章新增「深度学习与PyTorch」小节，并更新本章小结；
  2) 表4-1 算法分类统计表新增「深度学习」类别行；
  3) 第5章新增「深度学习算法模块」小节（含表5-8、图5-4/5-5/5-6）；
  4) 第6章新增深度学习功能测试（表6-3）；
  5) 第7章更新工作总结、创新点与不足；
  6) 更新图清单、表清单、变量注释表与摘要数量。

输出到 _最终.docx。
"""
import sys

sys.stdout.reconfigure(encoding="utf-8")

from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement

DOC_IN = "毕设/（先用1）科文学院本科生毕业论文模板-中文模板（理工科)_完善.docx"
DOC_OUT = "毕设/（先用1）科文学院本科生毕业论文模板-中文模板（理工科)_最终.docx"
FIG_DIR = "figures"

# ==================== 基础工具 ====================

def find_para(doc, substr):
    for p in doc.paragraphs:
        if substr in p.text:
            return p
    raise RuntimeError(f"未找到锚点: {substr!r}")

def insert_paragraph_after(paragraph):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    from docx.text.paragraph import Paragraph
    return Paragraph(new_p, paragraph._parent)

def _fmt_body(para):
    """正文格式：Normal、首行缩进24pt、1.5倍行距。"""
    pf = para.paragraph_format
    pf.first_line_indent = Pt(24)
    pf.line_spacing = 1.5

def add_heading_after(anchor, text):
    """插入二级标题（继承『论文二级标题』样式）。"""
    p = insert_paragraph_after(anchor)
    p.style = doc.styles["论文二级标题"]
    p.add_run(text)
    return p

def add_body_after(anchor, text):
    """插入正文段落（Normal、首行缩进、1.5行距）。"""
    p = insert_paragraph_after(anchor)
    p.add_run(text)
    _fmt_body(p)
    return p


def add_caption_after(anchor, text):
    """插入居中题注段落（表题/图题）。"""
    p = insert_paragraph_after(anchor)
    p.add_run(text)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.5
    return p


def add_picture_after(anchor, img_path, width_cm=14.0):
    """插入居中图片段落。"""
    p = insert_paragraph_after(anchor)
    p.add_run().add_picture(img_path, width=Cm(width_cm))
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def _set_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold


def make_table_after(anchor, header, rows, style="Table Grid"):
    tbl = doc.add_table(rows=len(rows) + 1, cols=len(header))
    tbl.style = doc.styles[style]
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    anchor._p.addnext(tbl._tbl)
    for j, h in enumerate(header):
        _set_text(tbl.rows[0].cells[j], h)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            _set_text(tbl.rows[i + 1].cells[j], val)
    _set_col_widths(tbl)
    return tbl


def _set_col_widths(tbl):
    ncol = len(tbl.columns)
    widths = {3: [Cm(3.4), Cm(4.2), Cm(6.4)], 4: [Cm(2.8), Cm(3.4), Cm(3.4), Cm(4.4)]}
    ws = widths.get(ncol, [Cm(14.0 / ncol)] * ncol)
    for row in tbl.rows:
        for j, w in enumerate(ws):
            row.cells[j].width = w


def replace_all(doc, mapping):
    """按顺序对全文段落执行字符串替换（保留首 run 格式）。"""
    n = 0
    for p in doc.paragraphs:
        for old, new in mapping:
            if old in p.text:
                fmt_run = p.runs[0] if p.runs else None
                txt = p.text
                for old2, new2 in mapping:
                    txt = txt.replace(old2, new2)
                for r in list(p.runs):
                    r._element.getparent().remove(r._element)
                run = p.add_run(txt)
                if fmt_run is not None:
                    run.font.name = fmt_run.font.name
                    run.font.size = fmt_run.font.size
                    run.font.bold = fmt_run.font.bold
                n += 1
                break
    print(f"  字符串替换命中 {n} 段")


# ==================== 数据定义 ====================

# 第2章新增小节内容
CH2_TITLE = "深度学习与PyTorch"
CH2_PARAS = [
    "深度学习是机器学习的一个重要分支，通过多层神经网络自动从数据中学习层次化的特征表示。在图像处理领域，卷积神经网络（CNN）是最具代表性的深度学习模型，其核心思想是利用局部感受野与权值共享，通过卷积运算自动提取从低级边缘纹理到高级语义概念的层级特征。典型的CNN由卷积层、池化层与全连接层堆叠而成：卷积层通过可学习的卷积核提取局部特征，池化层通过下采样降低特征图分辨率并增强平移不变性，全连接层则将提取的高层特征映射到最终的输出类别。",
    "本系统采用PyTorch作为深度学习框架。PyTorch是一个基于动态计算图的开源深度学习库，具有张量运算、自动微分与GPU加速三大核心能力。其动态图机制使得网络结构能够在运行时灵活构建与调试，便于教学演示中实时观察前向传播过程；torchvision作为PyTorch的计算机视觉工具库，提供了MobileNetV2、DeepLabV3等经典网络的结构定义及其在ImageNet、COCO等大规模数据集上预训练好的权重参数。",
    "预训练模型与迁移学习是本系统深度学习模块的重要设计基础。在ImageNet等大规模数据集上预训练的网络已经学习到丰富的通用视觉特征，将其作为特征提取器迁移到新的任务上，能够在小样本甚至零样本情况下取得较好效果。本系统采用的MobileNetV2是一种面向移动端设计的轻量级卷积网络，通过深度可分离卷积大幅降低参数量与计算量；DeepLabV3则通过空洞卷积与空间金字塔池化实现多尺度上下文信息的融合，是语义分割领域的代表性网络。",
    "基于上述技术，本系统构建了涵盖图像分类、语义分割、图像超分辨率与自定义图像分类四类任务的深度学习算法模块，并配套实现了模型训练流水线，使学生不仅能够使用预训练模型进行推理，还能亲手训练一个轻量级卷积网络，从而更全面地理解深度学习的完整工作流程。",
]
CH2_SUMMARY_ADD = "在经典图像处理技术之外，本章还介绍了卷积神经网络与PyTorch深度学习框架的基本原理，为系统深度学习算法模块的设计与实现奠定了理论基础。"

# 第5章新增小节内容
CH5_TITLE = "深度学习算法模块"
CH5_OVERVIEW = "为拓展系统的算法覆盖范围，本系统在经典图像处理算法之外，基于PyTorch实现了四类深度学习图像处理算法：图像分类、语义分割、图像超分辨率与CIFAR-10图像分类。与经典算法相同，四类深度学习算法遵循统一的函数签名（输入为BGR格式的图像数组，输出为处理结果图像），并在算法注册表中与经典算法统一管理。考虑到深度学习模型加载与推理的计算开销较大，系统通过Streamlit的缓存机制使模型在首次加载后常驻内存，避免每次交互都重新加载权重，从而在保证响应速度的同时降低计算资源的浪费。"
CH5_CLS = "图像分类算法采用MobileNetV2网络，该网络在ImageNet数据集上预训练，能够识别一千个类别。算法首先将输入图像缩放到224×224并按照ImageNet的均值和标准差进行归一化，然后送入网络前向传播得到各类别的得分，最后通过softmax函数将得分转换为概率分布，并选取概率最高的前k个类别连同其置信度以柱状图的形式叠加在结果图像下方（如图5-4(a)所示）。这种Top-k的展示方式有助于学生理解分类网络输出概率分布的物理含义。"
CH5_SEG = "语义分割算法采用DeepLabV3网络，该网络在COCO数据集上预训练，能够对图像中的每个像素赋予一个语义类别标签。算法将输入图像缩放到520×520并进行归一化后送入网络，得到每个像素在各类别上的得分，再通过argmax操作确定每个像素的预测类别，最后使用预先定义的调色板将类别索引映射为彩色掩膜并与原始图像进行加权融合，从而得到直观的分割结果（如图5-4(b)所示）。"
CH5_SR = "图像超分辨率算法采用SRCNN网络，旨在将低分辨率图像重建为高分辨率图像。SRCNN采用端到端的浅层卷积结构，直接学习低分辨率图像到高分辨率图像的非线性映射，是深度学习超分辨率的开创性工作之一。算法首先将低分辨率输入通过双三次插值放大到目标尺寸，再送入网络进行细节增强。相比传统的双三次插值方法，SRCNN能够重建出更加锐利的边缘与纹理细节（如图5-5(a)所示）。由于SRCNN需要针对特定任务训练，系统提供了相应的训练脚本，训练完成后即可在平台中进行推理。"
CH5_CIFAR = "CIFAR-10图像分类算法基于系统自定义设计的CIFAR10Net网络，用于识别CIFAR-10数据集中的十个物体类别。CIFAR10Net是一个包含若干卷积层与全连接层的小型卷积网络，专为32×32分辨率的小尺寸图像设计。该算法与前文所述基于预训练模型的图像分类算法的区别在于，其网络权重由系统提供的训练脚本在CIFAR-10数据集上从零训练得到（如图5-5(b)所示），完整展示了从数据准备、模型训练到推理部署的深度学习全流程。"
CH5_TRAIN = "系统在modules/training模块中提供了完整的训练流水线，支持CIFAR10Net与SRCNN两个模型的训练。训练流水线涵盖数据加载与增强、模型构建、损失计算、反向传播与参数更新、验证评估以及模型保存等环节：CIFAR-10数据集的训练图像经过随机裁剪与随机水平翻转进行数据增强，采用交叉熵损失函数与Adam优化器进行训练；SRCNN则以CIFAR-10图像下采样构造低分辨率与高分辨率训练样本对，采用均方误差损失函数进行训练。图5-6展示了两个模型训练过程中的损失与准确率变化曲线，从中可以看出随着训练轮次的增加，模型损失逐步下降、准确率稳步上升，验证了训练流水线的有效性。深度学习算法的主要参数如表5-8所示。"

# 表5-8 深度学习算法参数表
TABLE_5_8 = {
    "caption": "表5-8  深度学习算法参数表",
    "header": ["算法", "参数", "含义"],
    "rows": [
        ["图像分类 (MobileNetV2)", "top_k", "返回置信度最高的前k个类别"],
        ["语义分割 (DeepLabV3)", "opacity", "分割掩膜与原始图像的融合透明度"],
        ["超分辨率 (SRCNN)", "scale", "超分辨率放大倍数"],
        ["CIFAR-10分类 (CIFAR10Net)", "top_k", "返回置信度最高的前k个类别"],
    ],
}

# 图5-4/5-5/5-6
DL_FIGURES = [
    ("图5-4", "深度学习图像理解算法处理效果", "fig5_4_深度学习图像理解.png"),
    ("图5-5", "超分辨率重建与CIFAR-10分类效果", "fig5_5_超分与分类.png"),
    ("图5-6", "深度学习模型训练过程曲线", "fig5_6_训练曲线.png"),
]

# 表4-1 新增深度学习行
CATEGORY_DL_ROW = ["深度学习", "图像分类、语义分割、超分辨率、CIFAR-10分类", "4"]

# 第6章 深度学习功能测试
CH6_DL_TEST_INTRO = "针对四类深度学习算法，系统同样进行了功能测试，主要测试用例及结果如表6-3所示。"
TABLE_6_3 = {
    "caption": "表6-3  深度学习算法功能测试表",
    "header": ["测试项", "测试内容", "预期结果", "实际结果"],
    "rows": [
        ["图像分类", "输入示例图像，返回Top-3类别", "输出前3个类别及置信度", "与预期一致"],
        ["语义分割", "输入示例图像，进行像素级分割", "输出彩色分割掩膜", "与预期一致"],
        ["超分辨率", "输入低分辨率图像，3倍放大", "输出细节增强的高分辨率图像", "与预期一致"],
        ["CIFAR-10分类", "输入CIFAR-10测试图像", "输出Top-3类别及置信度", "与预期一致"],
    ],
}

# 第7章更新
WORK_SUMMARY_6 = "（6）基于PyTorch实现了图像分类、语义分割、图像超分辨率与CIFAR-10分类四类深度学习算法，并构建了配套的模型训练流水线，实现了从数据准备、模型训练到推理部署的完整深度学习流程。"
INNOVATION_ADD = "四是引入了深度学习算法模块，将预训练模型（MobileNetV2、DeepLabV3）与自定义训练的轻量级网络（SRCNN、CIFAR10Net）融入图像处理教学，使学生能够在同一平台上完成经典算法与深度学习算法的对比学习，拓展了平台的技术边界。"

# 变量注释表新增深度学习变量
DL_VARIABLES = [
    ["η", "学习率", "无量纲", "训练中控制参数更新步长的超参数"],
    ["epochs", "训练轮次", "次", "完整遍历训练集一遍的次数"],
    ["batch_size", "批大小", "张", "每次参数更新所使用的样本数量"],
    ["top_k", "Top-k", "个", "返回置信度最高的前k个类别"],
]

# 图清单/表清单最终条目
FINAL_FIGURES = [
    ("图2-1", "数字图像矩阵表示与二维卷积运算示意图"),
    ("图3-1", "系统总体用例图"),
    ("图4-1", "系统四层架构图"),
    ("图4-2", "系统数据流图"),
    ("图4-3", "系统界面布局示意图"),
    ("图5-1", "系统内置示例图像"),
    ("图5-2", "滤波与图像增强算法处理效果对比"),
    ("图5-3", "边缘检测与形态学操作处理效果对比"),
    ("图5-4", "深度学习图像理解算法处理效果"),
    ("图5-5", "超分辨率重建与CIFAR-10分类效果"),
    ("图5-6", "深度学习模型训练过程曲线"),
    ("图6-1", "各算法平均处理耗时对比图"),
    ("图6-2", "处理耗时随图像分辨率变化曲线"),
]
FINAL_TABLES = [
    ("表4-1", "算法分类统计表"),
    ("表5-1", "开发环境配置表"),
    ("表5-2", "滤波算法参数表"),
    ("表5-3", "图像增强算法参数表"),
    ("表5-4", "边缘检测算法参数表"),
    ("表5-5", "形态学操作参数表"),
    ("表5-6", "阈值分割算法参数表"),
    ("表5-7", "图像变换算法参数表"),
    ("表5-8", "深度学习算法参数表"),
    ("表6-1", "功能测试用例表"),
    ("表6-2", "部分算法平均处理耗时（800×600图像）"),
    ("表6-3", "深度学习算法功能测试表"),
]


def fill_list(table, entries):
    n_data = len(entries)
    need = n_data + 1
    while len(table.rows) < need:
        table.add_row()
    while len(table.rows) > need:
        table._tbl.remove(table.rows[-1]._tr)
    for i, (num, name) in enumerate(entries):
        row = table.rows[i + 1]
        _set_text(row.cells[0], num)
        _set_text(row.cells[1], name)
        _set_text(row.cells[2], "")


# ==================== 主流程 ====================

def main():
    global doc
    doc = Document(DOC_IN)

    # ---- 1. 第2章：深度学习与PyTorch ----
    print("== 第2章：插入深度学习小节 ==")
    ch2_anchor = find_para(doc, "体现了Python科学计算生态的成熟与高效")
    p = ch2_anchor
    p = add_heading_after(p, CH2_TITLE)
    for para in CH2_PARAS:
        p = add_body_after(p, para)
    print("  已插入小节正文")

    # 更新本章小结
    summary_anchor = find_para(doc, "上述技术的合理选型与分工")
    summary_anchor.add_run(CH2_SUMMARY_ADD)
    print("  已更新本章小结")

    # ---- 2. 表4-1 新增深度学习行 ----
    print("== 表4-1：新增深度学习类别行 ==")
    tbl_cat = doc.tables[3]
    tbl_cat.add_row()
    last = tbl_cat.rows[-1]
    for j, val in enumerate(CATEGORY_DL_ROW):
        _set_text(last.cells[j], val)
    print("  已新增行:", CATEGORY_DL_ROW)

    # ---- 3. 第5章：深度学习算法模块 ----
    print("== 第5章：插入深度学习算法模块 ==")
    ch5_anchor = find_para(doc, "图像变换算法参数表")
    p = ch5_anchor
    p = add_heading_after(p, CH5_TITLE)
    p = add_body_after(p, CH5_OVERVIEW)
    p = add_body_after(p, CH5_CLS)
    p = add_body_after(p, CH5_SEG)
    p = add_body_after(p, CH5_SR)
    p = add_body_after(p, CH5_CIFAR)
    p = add_body_after(p, CH5_TRAIN)

    # 表5-8
    cap = add_caption_after(p, TABLE_5_8["caption"])
    tbl = make_table_after(cap, TABLE_5_8["header"], TABLE_5_8["rows"])
    # 图5-4/5-5/5-6：依次插入在表格之后
    from docx.text.paragraph import Paragraph
    parent_proxy = cap._parent
    anchor_el = tbl._tbl
    for (figno, figname, fname) in DL_FIGURES:
        pic_el = OxmlElement("w:p")
        anchor_el.addnext(pic_el)
        pp = Paragraph(pic_el, parent_proxy)
        pp.add_run().add_picture(f"{FIG_DIR}/{fname}", width=Cm(14.0))
        pp.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap_el = OxmlElement("w:p")
        pic_el.addnext(cap_el)
        cp = Paragraph(cap_el, parent_proxy)
        cp.add_run(f"{figno}  {figname}")
        cp.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.line_spacing = 1.5
        anchor_el = cap_el
    print("  已插入 DL 小节、表5-8 与三幅图")

    # ---- 4. 第6章：深度学习功能测试 ----
    print("== 第6章：插入深度学习功能测试 ==")
    ch6_anchor = find_para(doc, "30种算法均能正确处理图像")  # [262]，将被替换前定位
    p = add_body_after(ch6_anchor, CH6_DL_TEST_INTRO)
    cap = add_caption_after(p, TABLE_6_3["caption"])
    make_table_after(cap, TABLE_6_3["header"], TABLE_6_3["rows"])
    print("  已插入 表6-3")

    # ---- 5. 第7章：更新总结 ----
    print("== 第7章：更新总结 ==")
    # 新增工作总结第(6)条
    wp5 = find_para(doc, "保障了系统进行图像处理的响应速度和稳定性")  # (5)
    add_body_after(wp5, WORK_SUMMARY_6)
    # 创新点追加
    inno = find_para(doc, "相比传统桌面端工具具有部署便捷")
    inno.add_run(INNOVATION_ADD)
    print("  已新增总结第(6)条与创新点")

    # ---- 6. 全局数量替换 ----
    print("== 全局数量替换 ==")
    mapping = [
        ("六大类共30余种算法", "七大类共34种算法"),
        ("六大类30种算法", "七大类34种算法"),
        ("六大类共30种", "七大类共34种"),
        ("30种算法", "34种算法"),
        ("约两千行", "约两千余行"),
    ]
    replace_all(doc, mapping)

    # 修正「未涉及深度学习」的不足表述
    print("== 修正不足与展望 ==")
    old_lim = "其一，系统的算法覆盖主要集中在经典算法层面，未涉及深度学习相关的图像处理算法，后续可引入卷积神经网络等深度学习模型，丰富算法体系；"
    new_lim = "其一，深度学习算法的运行依赖较大规模的计算资源，在无GPU的普通计算机上推理速度较慢，后续可考虑引入模型量化、剪枝等加速技术；"
    for p in doc.paragraphs:
        if old_lim in p.text:
            fmt_run = p.runs[0] if p.runs else None
            txt = p.text.replace(old_lim, new_lim)
            for r in list(p.runs):
                r._element.getparent().remove(r._element)
            run = p.add_run(txt)
            if fmt_run is not None:
                run.font.name = fmt_run.font.name
                run.font.size = fmt_run.font.size
                run.font.bold = fmt_run.font.bold
            print("  已修正不足表述")
            break

    # ---- 6.5 修正分类口径表述 ----
    print("== 修正分类口径表述 ==")
    phrasing = [
        ("将常用算法划分为滤波、图像增强、边缘检测、形态学操作、阈值分割、图像变换六大类，基于OpenCV实现共34种算法的处理功能。",
         "将常用算法划分为滤波、图像增强、边缘检测、形态学操作、阈值分割、图像变换等六大类经典算法与一类深度学习算法，基于OpenCV与PyTorch实现共34种算法的处理功能。"),
        ("系统将34种算法划分为六个类别",
         "系统将34种算法划分为七个类别"),
        ("覆盖滤波、图像增强、边缘检测、形态学操作、阈值分割、图像变换七大类共34种",
         "覆盖滤波、图像增强、边缘检测、形态学操作、阈值分割、图像变换六大类经典算法与四类深度学习算法，共34种"),
    ]
    for old, new in phrasing:
        for p in doc.paragraphs:
            if old in p.text:
                fmt_run = p.runs[0] if p.runs else None
                txt = p.text.replace(old, new)
                for r in list(p.runs):
                    r._element.getparent().remove(r._element)
                run = p.add_run(txt)
                if fmt_run is not None:
                    run.font.name = fmt_run.font.name
                    run.font.size = fmt_run.font.size
                    run.font.bold = fmt_run.font.bold
                print(f"  已修正: {new[:24]}...")
                break

    # ---- 7. 图清单 / 表清单 ----
    print("== 重建图清单 / 表清单 ==")
    fill_list(doc.tables[0], FINAL_FIGURES)
    fill_list(doc.tables[1], FINAL_TABLES)
    print(f"  图清单 {len(FINAL_FIGURES)} 条，表清单 {len(FINAL_TABLES)} 条")

    # ---- 8. 变量注释表：新增 DL 变量 ----
    print("== 变量注释表：新增 DL 变量 ==")
    vtab = doc.tables[2]
    for var in DL_VARIABLES:
        vtab.add_row()
        for j, val in enumerate(var):
            _set_text(vtab.rows[-1].cells[j], val)
    print("  已新增 4 个 DL 变量")

    # ---- 9. 摘要数量更新 ----
    print("== 更新摘要数量 ==")
    pairs = [
        ("该论文有图10幅，表10个，参考文献160篇", "该论文有图13幅，表12个，参考文献160篇"),
        ("This paper contains 10 figures, 10 tables, and 160 references",
         "This paper contains 13 figures, 12 tables, and 160 references"),
    ]
    for old, new in pairs:
        for p in doc.paragraphs:
            if old in p.text:
                fmt_run = p.runs[0] if p.runs else None
                for r in list(p.runs):
                    r._element.getparent().remove(r._element)
                run = p.add_run(new)
                if fmt_run is not None:
                    run.font.name = fmt_run.font.name
                    run.font.size = fmt_run.font.size
                    run.font.bold = fmt_run.font.bold
                print(f"  已更新: {new[:30]}...")
                break

    doc.save(DOC_OUT)
    print(f"\n已保存: {DOC_OUT}")


if __name__ == "__main__":
    main()
