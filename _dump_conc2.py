# -*- coding: utf-8 -*-
import sys
sys.stdout.reconfigure(encoding="utf-8")
from docx import Document

P = r'd:\1\bishe\毕设\7230267148_杨康康_基于opencv和streamlit的图像处理算法交互式演示系统 - 降重改写版.docx'
doc = Document(P)
ps = doc.paragraphs
print("===== 320-345 =====")
for i in range(320, min(345, len(ps))):
    t = ps[i].text.strip()
    if t:
        print(f"[{i}] <{ps[i].style.name}> {t[:130]}")
print("===== 表格清单 (tables[1]) =====")
for ri, row in enumerate(doc.tables[1].rows):
    print(ri, " | ".join(c.text.strip() for c in row.cells))
print("===== 表6-4/6-5 正文引用段 =====")
for i in (283, 285, 288, 291, 295, 298, 301, 305, 306):
    print(f"[{i}] {ps[i].text.strip()[:200]}")
print("===== 表格总数 =====", len(doc.tables))
