# -*- coding: utf-8 -*-
"""统一 300ms 指标口径:仅约束算法核心(本体)计算耗时
- 3.3 指标改述
- 摘要(中/英)口径同步
- 6.4 结果分析结论口径同步
- 附:6.1/6.3/6.5 中仍用"实时性/交互响应"的表述一并对齐
"""
from docx import Document

PATH = r'd:\1\bishe\毕设\7230267148_杨康康_基于opencv和streamlit的图像处理算法交互式演示系统 - 降重改写版.docx'

EDITS = [
    # 3.3 非功能需求(性能) —— 改述为仅约束算法核心耗时
    ("本系统将算法核心单次执行的计算耗时控制在300毫秒以内，为用户调参后快速获得处理结果提供了保障；",
     "本系统的实时性指标只约束算法本体的计算耗时，即算法核心单次执行的计算耗时需要控制在300毫秒以内；"),
    # 中文摘要 —— 同步口径
    ("低于300ms的实时性指标算是达标",
     "低于300ms这一算法核心计算耗时的指标要求，算是达标"),
    # 英文摘要 —— 同步口径并统一数值
    ("Tests show that the algorithm takes about 3ms to process in average time, meeting real-time performance indicators below 300ms.",
     "Tests show that the twelve representative algorithms take about 3.7 ms on average (median about 1.8 ms) in algorithm core computation, meeting the requirement that the algorithm core computation time stays below 300 ms."),
    # 6.4 结果分析 —— 同步口径
    ("算法处理耗时远低于300毫秒这一算法计算环节的实时性约束",
     "算法本体的计算耗时远低于300毫秒这一算法核心耗时的约束"),
    # 6.1 测试方法 —— 去掉"交互响应"口径
    ("看各个算法能不能达到300ms的交互响应目标",
     "看各个算法本体的计算耗时能不能达到300ms这一指标要求"),
    # 6.3 性能测试 —— 去掉端到端"系统响应延迟"表述
    ("跟300毫秒的实时性指标要求比起来要低得多，系统响应延迟的主要来源是算法本身的处理时间",
     "跟300毫秒这一算法核心计算耗时的指标要求比起来要低得多"),
    # 6.3 性能测试 —— "实时性红线"改为算法核心耗时约束
    ("也都没有越过300毫秒这条实时性红线",
     "也都没有越过300毫秒这条算法核心耗时的约束"),
    # 6.5 本章小结 —— 同步口径
    ("离300毫秒这条实时性还有打量余量",
     "离300毫秒这条算法核心耗时的约束还有打量余量"),
]


def set_text(p, text):
    if not p.runs:
        p.add_run(text)
        return
    p.runs[0].text = text
    for r in list(p.runs[1:]):
        r._element.getparent().remove(r._element)


doc = Document(PATH)
done, missed = [], list(EDITS)
for p in doc.paragraphs:
    t = p.text
    changed = False
    for pair in list(missed):
        old, new = pair
        if old in t:
            t = t.replace(old, new)
            done.append(old[:16])
            missed.remove(pair)
            changed = True
    if changed:
        set_text(p, t)

doc.save(PATH)
print(f"已应用 {len(done)} 处:")
for d in done:
    print("  -", d)
print(f"未命中 {len(missed)} 处:")
for old, _ in missed:
    print("  !", old[:24])
