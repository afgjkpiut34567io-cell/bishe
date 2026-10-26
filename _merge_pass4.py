# -*- coding: utf-8 -*-
"""合并pass4结果,核验,应用生成最终文件"""
from docx import Document
import json, glob, re

SRC = r'd:\1\bishe\毕设\7230267148_杨康康_基于opencv和streamlit的图像处理算法交互式演示系统 - 降重改写版.docx'
DST = r'd:\1\bishe\毕设\7230267148_杨康康_基于opencv和streamlit的图像处理算法交互式演示系统 - 降重改写版2.docx'

doc = Document(SRC)
paras = doc.paragraphs

# 载入pass4输入(原文)与输出(新文)
ins = {}
for f in sorted(glob.glob(r'd:\1\bishe\_rew\pass4_in*.json')):
    for it in json.load(open(f, encoding='utf-8')):
        ins[str(it['idx'])] = it['text']
outs = {}
for f in sorted(glob.glob(r'd:\1\bishe\_rew\pass4_out*.json')):
    outs.update(json.load(open(f, encoding='utf-8')))

print(f"输入段数: {len(ins)}, 输出段数: {len(outs)}")
missing = [k for k in ins if k not in outs]
print(f"缺失: {missing}")

# 核验
FORBIDDEN = ['说白了', '可以这么讲', '不难发现', '值得注意', '综上所述', '总而言之', '总的来说',
             '换言之', '由此可见', '我觉得', '我认为', '在我看来']
bad_len, bad_tok, bad_word = [], [], []
tot_o, tot_n = 0, 0
for k, o in ins.items():
    n = outs.get(k, '')
    tot_o += len(o)
    tot_n += len(n)
    r = len(n) / max(len(o), 1)
    if not (0.8 <= r <= 1.2):
        bad_len.append((k, len(o), len(n), round(r, 2)))
    for pat in [r'\[\d+\]', r'图\d+(?:-\d+)?', r'表\d+(?:-\d+)?']:
        if sorted(re.findall(pat, o)) != sorted(re.findall(pat, n)):
            bad_tok.append((k, pat, re.findall(pat, o), re.findall(pat, n)))
    for w in FORBIDDEN:
        if w in n:
            bad_word.append((k, w))

print(f"总字数: {tot_o} -> {tot_n} ({tot_n/max(tot_o,1):.3f})")
print(f"字数越界: {bad_len}")
print(f"引用图表不一致: {bad_tok}")
print(f"红线词: {bad_word}")

# 应用
applied = 0
for k, v in outs.items():
    p = paras[int(k)]
    if not p.runs:
        print(f"跳过无run: {k}")
        continue
    p.runs[0].text = v
    for r in list(p.runs[1:]):
        r._element.getparent().remove(r._element)
    applied += 1
doc.save(DST)
print(f"已应用{applied}段, 保存: {DST}")

# 全文总字数
def total(path):
    d = Document(path)
    c = sum(len(re.sub(r'\s', '', p.text)) for p in d.paragraphs)
    t = sum(len(re.sub(r'\s', '', c2.text)) for tb in d.tables for row in tb.rows for c2 in row.cells)
    return c + t
print(f"原改写版总字数 {total(SRC)}, 新文件总字数 {total(DST)}")
