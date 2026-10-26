# -*- coding: utf-8 -*-
"""为23个偏移标红段提取正确文本(探针映射到文档真实段落)"""
from docx import Document
import json

rep = Document(r'd:\download\免费_Word标红版_AIGC检测报告_[7230267148_杨康康_基于ope] (2).docx')
doc = Document(r'd:\1\bishe\毕设\7230267148_杨康康_基于opencv和streamlit的图像处理算法交互式演示系统 - 降重改写版2.docx')

mis_rids = [209, 232, 260, 263, 277, 278, 296, 297, 298, 300, 301, 308, 309, 319, 323, 329, 331, 332, 343, 344, 345, 346, 348]
rep_texts = [p.text.strip() for p in rep.paragraphs]
doc_texts = [p.text.strip() for p in doc.paragraphs]

items = []
for rid in mis_rids:
    probe = rep_texts[rid][:20]
    found = None
    for di, dt in enumerate(doc_texts):
        if dt and dt[:20] == probe:
            found = di
            break
    assert found is not None, f"未找到 {rid}"
    items.append({'rid': rid, 'idx': found, 'text': doc_texts[found]})

print(f"23段总字数: {sum(len(i['text']) for i in items)}")
json.dump(items[:12], open(r'd:\1\bishe\_rew\pass5b_in0.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
json.dump(items[12:], open(r'd:\1\bishe\_rew\pass5b_in1.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
print("批0:", len(items[:12]), "段, idx", [i['idx'] for i in items[:12]])
print("批1:", len(items[12:]), "段, idx", [i['idx'] for i in items[12:]])
