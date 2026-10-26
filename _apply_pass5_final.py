# -*- coding: utf-8 -*-
"""合成最终版: pass5(19段直接匹配) + pass5b(23段偏移) -> 降重改写版2 => 降重改写版3"""
from docx import Document
import json, glob, re

SRC = r'd:\1\bishe\毕设\7230267148_杨康康_基于opencv和streamlit的图像处理算法交互式演示系统 - 降重改写版2.docx'
DST = r'd:\1\bishe\毕设\7230267148_杨康康_基于opencv和streamlit的图像处理算法交互式演示系统 - 降重改写版3.docx'

doc = Document(SRC)
paras = doc.paragraphs

# 19个直接匹配段: doc idx == 报告rid
identity_idx = [51, 82, 84, 91, 92, 93, 112, 117, 120, 121, 123, 125, 132, 138, 181, 182, 183, 185, 186]

# pass5输出(键为rid, 直接匹配段rid==doc idx)
p5 = {}
for f in sorted(glob.glob(r'd:\1\bishe\_rew\pass5_out*.json')):
    p5.update(json.load(open(f, encoding='utf-8')))
# pass5b输出(键为doc idx)
p5b = {}
for f in sorted(glob.glob(r'd:\1\bishe\_rew\pass5b_out*.json')):
    p5b.update(json.load(open(f, encoding='utf-8')))

# 合成应用表: {doc_idx: new_text}
apply_map = {}
for i in identity_idx:
    k = str(i)
    if k in p5:
        apply_map[i] = p5[k]
    else:
        print(f"!! pass5缺identity段 {k}")
for k, v in p5b.items():
    apply_map[int(k)] = v
print(f"应用段数: {len(apply_map)} (应为42)")

# 核验
FORBID_SUBSTR = ['我觉得', '我认为', '我选', '我做', '我在', '我把', '我用', '我按', '我靠', '我又', '让我', '给我们']
bad = []
tot_o, tot_n = 0, 0
for did, v in sorted(apply_map.items()):
    o = paras[did].text.strip()
    tot_o += len(o)
    tot_n += len(v)
    r = len(v) / max(len(o), 1)
    if not (0.8 <= r <= 1.2):
        bad.append(('字数', did, len(o), len(v), round(r, 2)))
    for w in FORBID_SUBSTR:
        if w in v:
            bad.append(('人称', did, w))
    for pat in [r'\[\d+\]', r'图\d+(?:-\d+)?', r'表\d+(?:-\d+)?']:
        if sorted(re.findall(pat, o)) != sorted(re.findall(pat, v)):
            bad.append(('标记', did, pat))
print(f"字数: {tot_o} -> {tot_n}")
print("问题:", bad if bad else "无")

# 应用
applied = 0
for did, v in apply_map.items():
    p = paras[did]
    if not p.runs:
        print(f"跳过无run: {did}")
        continue
    p.runs[0].text = v
    for r in list(p.runs[1:]):
        r._element.getparent().remove(r._element)
    applied += 1
doc.save(DST)
print(f"已应用{applied}段, 保存: {DST}")

def total(path):
    d = Document(path)
    c = sum(len(re.sub(r'\s', '', p.text)) for p in d.paragraphs)
    t = sum(len(re.sub(r'\s', '', c2.text)) for tb in d.tables for row in tb.rows for c2 in row.cells)
    return c + t
print(f"源总字数 {total(SRC)}, 新总字数 {total(DST)}")
