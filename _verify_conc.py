# -*- coding: utf-8 -*-
import sys
sys.stdout.reconfigure(encoding="utf-8")
from docx import Document
P = r'd:\1\bishe\毕设\7230267148_杨康康_基于opencv和streamlit的图像处理算法交互式演示系统 - 降重改写版.docx'
doc = Document(P)
ps = doc.paragraphs
print("段落总数", len(ps), "表格总数", len(doc.tables))
KEYS = ["从技术可行性", "并发规模和部署边界", "会话状态数据始终保持着一致",
        "并发访问的实测", "表6-4 并发实测", "从表6-4的数据", "表6-5 超分辨率",
        "表6-6 加噪", "本系统的创新点"]
for i, p in enumerate(ps):
    t = p.text.strip()
    for k in KEYS:
        if k in t:
            print(f"\n[{i}] <{p.style.name}> {t}")
            break
print("\n===== 表清单 =====")
for row in doc.tables[1].rows:
    print("  ", [c.text.strip() for c in row.cells])
