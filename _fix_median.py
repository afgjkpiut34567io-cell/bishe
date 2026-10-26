# -*- coding: utf-8 -*-
"""[审查项#19] 统一耗时统计量为中位数：
  6.1 测试方法「取中位数」；表6-3 表头/题注、图6-1 题注与图内轴标签、6.3 正文一律改称中位数耗时；
  同步图清单/表清单条目，并重新生成并回嵌图6-1 插图。"""
import sys, os, shutil
sys.stdout.reconfigure(encoding="utf-8")
from docx import Document
from matplotlib import font_manager

PATH = r'd:\1\bishe\毕设\7230267148_杨康康_基于opencv和streamlit的图像处理算法交互式演示系统 - 降重改写版.docx'
BAK = r'd:\1\bishe\毕设\_bak_降重改写版_20261014_before_median.docx'
FIG = r'd:\1\bishe\figures\fig6_1_耗时对比图.png'
shutil.copy2(PATH, BAK)
shutil.copy2(FIG, r'd:\1\bishe\figures\_bak_fig6_1_耗时对比图.png')
print("备份:", BAK, "|", r'd:\1\bishe\figures\_bak_fig6_1_耗时对比图.png')

# ---------------- 1) 重新生成图6-1（仅改轴标签为“中位数耗时 (ms)”） ----------------
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

_font = None
for _n in ("SimHei", "Microsoft YaHei", "SimSun"):
    if any(f.name == _n for f in font_manager.fontManager.ttflist):
        _font = _n; break
plt.rcParams["font.sans-serif"] = [_font or "DejaVu Sans"]
plt.rcParams["axes.unicode_minus"] = False

data = [("傅里叶变换", 15.8), ("双边滤波", 12.5), ("Canny", 3.8),
        ("自适应阈值", 2.6), ("CLAHE", 2.3), ("Sobel", 2.1),
        ("中值滤波", 1.5), ("高斯滤波", 1.2), ("Otsu", 1.0),
        ("腐蚀", 0.9), ("直方图均衡", 0.8), ("伽马校正", 0.4)]
data.sort(key=lambda x: x[1])
names = [d[0] for d in data]; vals = [d[1] for d in data]
fig, ax = plt.subplots(figsize=(8.5, 4.6))
colors = ["#E08A3C" if v > 3 else "#2C5F8A" for v in vals]
bars = ax.barh(names, vals, color=colors, edgecolor="white")
ax.set_xlabel("中位数耗时 (ms)", fontsize=11)
ax.set_xlim(0, 18)
for b, v in zip(bars, vals):
    ax.text(v + 0.2, b.get_y() + b.get_height() / 2, f"{v:.1f}", va="center", fontsize=9)
ax.spines["top"].set_visible(False); ax.spines["right"].set_visible(False)
fig.tight_layout()
fig.savefig(FIG, dpi=200, bbox_inches="tight", facecolor="white")
plt.close(fig)
NEW_PNG = open(FIG, "rb").read()
print("新图:", os.path.getsize(FIG), "字节")

# ---------------- 2) 文档文本改动 ----------------
doc = Document(PATH)

p = doc.paragraphs[271]
assert "重复执行取平均值" in p.runs[0].text
p.runs[0].text = p.runs[0].text.replace("重复执行取平均值", "重复执行取中位数")

p = doc.paragraphs[284]
p.runs[0].text = p.runs[0].text.replace("平均处理耗时", "中位数耗时")
assert "平均" not in p.text

p = doc.paragraphs[285]
assert p.runs[2].text == "部分算法平均处理耗时（"
p.runs[2].text = "部分算法中位数耗时（"

p = doc.paragraphs[286]
p.runs[0].text = p.runs[0].text.replace("平均处理耗时", "中位数耗时")
assert "平均" not in p.text

p = doc.paragraphs[288]
assert p.runs[2].text == "各算法平均处理耗时对比图"
p.runs[2].text = "各算法中位数耗时对比图"

# 表6-3 表头（双列对布局，两处）
for c in doc.tables[16].rows[0].cells:
    if "平均耗时" in c.text:
        c.paragraphs[0].runs[0].text = "中位数耗时（"
print("表6-3 表头:", [c.text for c in doc.tables[16].rows[0].cells])

# 前置图清单 / 表清单
figlist, tablist = doc.tables[0], doc.tables[1]
for r in figlist.rows:
    if r.cells[0].text.strip() == "图6-1":
        r.cells[1].paragraphs[0].runs[0].text = "各算法中位数耗时对比图"
for r in tablist.rows:
    if r.cells[0].text.strip() == "表6-3":
        r.cells[1].paragraphs[0].runs[0].text = "部分算法中位数耗时（800×600图像）"

# ---------------- 3) 回嵌新图6-1 ----------------
part = doc.part.related_parts["rId25"]
old_len = len(part.blob)
assert old_len == 70087 and part.blob[:8] == b"\x89PNG\r\n\x1a\n", (old_len,)
part._blob = NEW_PNG
print("图6-1 回嵌: 旧", old_len, "字节 -> 新", len(part.blob), "字节")

doc.save(PATH)
print("\n已保存。校验：")
d2 = Document(PATH)
for n in (271, 284, 285, 286, 288):
    print(n, "|", d2.paragraphs[n].text)
print("表6-3 表头:", [c.text for c in d2.tables[16].rows[0].cells])
print("图清单 图6-1:", [r.cells[1].text for r in d2.tables[0].rows if r.cells[0].text.strip() == "图6-1"])
print("表清单 表6-3:", [r.cells[1].text for r in d2.tables[1].rows if r.cells[0].text.strip() == "表6-3"])
print("图6-1 blob 长度:", len(d2.part.related_parts["rId25"].blob))
