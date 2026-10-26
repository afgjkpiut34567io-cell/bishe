# -*- coding: utf-8 -*-
"""验证报告(2)标红段与文档的真实映射(文本探针)"""
from docx import Document
import json

rep = Document(r'd:\download\免费_Word标红版_AIGC检测报告_[7230267148_杨康康_基于ope] (2).docx')
doc = Document(r'd:\1\bishe\毕设\7230267148_杨康康_基于opencv和streamlit的图像处理算法交互式演示系统 - 降重改写版2.docx')

red = []
for i, p in enumerate(rep.paragraphs):
    t = p.text.strip()
    if not t:
        continue
    rc = sum(len(r.text) for r in p.runs if r.font.color and str(r.font.color.rgb) == 'F12828')
    total = sum(len(r.text) for r in p.runs)
    if total > 30 and rc / total > 0.3:
        red.append((i, t))

doc_texts = [p.text.strip() for p in doc.paragraphs]
mismatch = []
ok = 0
for rid, rtext in red:
    # 报告idx与文档idx是否直接一致
    if rid < len(doc_texts) and doc_texts[rid] and doc_texts[rid][:25] == rtext[:25]:
        ok += 1
        continue
    # 探针全文找
    found = None
    for di, dt in enumerate(doc_texts):
        if dt and dt[:20] == rtext[:20]:
            found = di
            break
    mismatch.append((rid, found, rtext[:45]))

print(f"直接一致: {ok}/{len(red)}")
print(f"不一致或需探针的: {len(mismatch)}")
for m in mismatch:
    print("  ", m)
