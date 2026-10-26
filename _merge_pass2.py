# -*- coding: utf-8 -*-
"""合并pass2结果,复核重合度,重新生成最终docx"""
from docx import Document
import json, glob, re, difflib

SRC = r'd:\1\bishe\毕设\7230267148_杨康康_基于opencv和streamlit的图像处理算法交互式演示系统.docx'
DST = r'd:\1\bishe\毕设\7230267148_杨康康_基于opencv和streamlit的图像处理算法交互式演示系统 - 降重改写版.docx'

orig = Document(SRC).paragraphs

# pass1 + pass2 合并(pass2覆盖pass1)
rew = {}
for f in sorted(glob.glob(r'd:\1\bishe\_rew\b*.json')):
    rew.update(json.load(open(f, encoding='utf-8')))
n1 = len(rew)
for f in sorted(glob.glob(r'd:\1\bishe\_rew\pass2_out*.json')):
    rew.update(json.load(open(f, encoding='utf-8')))
print(f"pass1段数: {n1}, 合并后: {len(rew)}")

# 复核重合度
def overlap13(a, b):
    grams = set(b[i:i+13] for i in range(len(b)-12))
    if not grams:
        return 0
    return sum(1 for i in range(len(a)-12) if a[i:i+13] in grams) / max(len(a)-12, 1)

def lcs(a, b):
    m = difflib.SequenceMatcher(None, a, b, autojunk=False)
    return max(m.get_matching_blocks(), key=lambda x: x.size).size

rates, worst = [], []
for k, v in rew.items():
    o = re.sub(r'\s', '', orig[int(k)].text)
    n = re.sub(r'\s', '', v)
    r = overlap13(o, n)
    rates.append(r)
    if r > 0.22:
        worst.append((int(k), round(r, 2), lcs(o, n)))
worst.sort()
import statistics
print(f"13gram重合率: 平均{statistics.mean(rates):.3f}, 最大{max(rates):.3f}")
print(f"仍>0.22的段: {len(worst)}")
for w in worst[:10]:
    print("  ", w)

# 引用/图表标记复核
bad_tok = []
for k, v in rew.items():
    o = orig[int(k)].text
    for pat in [r'\[\d+\]', r'图\d+(?:-\d+)?', r'表\d+(?:-\d+)?']:
        if sorted(re.findall(pat, o)) != sorted(re.findall(pat, v)):
            bad_tok.append((k, pat))
print(f"引用/图表标记不一致: {bad_tok}")

# 红线词复核
FORBIDDEN = ['说白了', '就拿', '可以这么讲', '不难发现', '值得注意', '综上所述', '总而言之',
             '总的来说', '换言之', '与此同时', '由此可见', '我觉得', '我认为', '在我看来']
hits = [(k, w) for k, v in rew.items() for w in FORBIDDEN if w in v]
print(f"红线词命中: {hits}")

# 应用到docx
doc = Document(SRC)
paras = doc.paragraphs
applied = 0
for k, v in rew.items():
    p = paras[int(k)]
    if not p.runs:
        continue
    p.runs[0].text = v
    for r in list(p.runs[1:]):
        r._element.getparent().remove(r._element)
    applied += 1
doc.save(DST)
print(f"已应用{applied}段, 已保存: {DST}")

tot_o = sum(len(re.sub(r'\s','',orig[int(k)].text)) for k in rew)
tot_n = sum(len(re.sub(r'\s','',v)) for k, v in rew.items())
print(f"改写段总字数: {tot_o} -> {tot_n} ({tot_n/tot_o:.3f})")
