# -*- coding: utf-8 -*-
"""
向 6.4 结果分析追加客观指标内容：文字段 + 表6-4 + 表6-5 + 图6-3，并同步前置图/表清单
"""
import sys, json
sys.stdout.reconfigure(encoding="utf-8")
from docx import Document
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

PATH = r'd:\1\bishe\毕设\7230267148_杨康康_基于opencv和streamlit的图像处理算法交互式演示系统 - 降重改写版.docx'
FIG = r'd:\1\bishe\figures\fig6_3_参数消融曲线.png'
M = json.load(open(r'd:\1\bishe\_ch6_metrics.json', encoding='utf-8'))

doc = Document(PATH)

# ---------- 工具 ----------
def para_after(el, parent):
    p_el = OxmlElement('w:p')
    el.addnext(p_el)
    return Paragraph(p_el, parent)

def body(el, parent, text):
    p = para_after(el, parent)
    p.add_run(text)
    p.paragraph_format.first_line_indent = Pt(24)
    p.paragraph_format.line_spacing = 1.5
    return p

def caption(el, parent, text):
    p = para_after(el, parent)
    p.add_run(text)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.5
    return p

def picture(el, parent, path, width_cm=13.5):
    p = para_after(el, parent)
    p.add_run().add_picture(path, width=Cm(width_cm))
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p

def table_after(el, header, rows, widths):
    tbl = doc.add_table(rows=len(rows) + 1, cols=len(header))
    tbl.style = doc.styles['Table Grid']
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    el.addnext(tbl._tbl)
    for j, h in enumerate(header):
        c = tbl.rows[0].cells[j]; c.text = ''
        cc = c.paragraphs[0]; cc.add_run(h); cc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = tbl.rows[i + 1].cells[j]; c.text = ''
            cc = c.paragraphs[0]; cc.add_run(str(v)); cc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in tbl.rows:
        for j, w in enumerate(widths):
            r.cells[j].width = Cm(w)
    return tbl

# ---------- 定位锚点：6.4 中"教学辅助的角度"段 ----------
anchor = None
for p in doc.paragraphs:
    if "教学辅助的角度" in p.text:
        anchor = p
        break
assert anchor is not None, "未找到 6.4 锚点"
parent = anchor._parent
cur = anchor._p

# ---------- 正文内容 ----------
p = body(cur, parent,
    "前面几节把功能、性能还有稳定性都测过一遍，结论大体有了，不过前面对算法处理效果的评价，主要还是停留在目视检查这一步，"
    "缺少客观数字来支撑。为此测试环节又补了一轮测量，用的是系统自带的客观质量指标模块。这个模块对处理前后的图像逐像素算出三个量："
    "MSE是均方误差，越小说明结果离开原图越近；PSNR叫峰值信噪比，单位是分贝，数值越大越接近原图；"
    "SSIM叫结构相似度，取值一般在0到1中间，越靠近1，图像的结构保持得越好。")
cur = p._p

p = body(cur, parent,
    "超分辨率这一块，拿内置示例图做了对照。先把原图的亮度分量按2倍和3倍下采样，得到低分辨率输入，"
    "再用两种路子把它重建回原尺寸，一条是最常见的双三次插值，另一条是系统里的SRCNN网络，"
    "最后都拿重建结果跟原始高分辨率图像按亮度分量逐像素比较，把PSNR和SSIM算出来。对照结果见表6-4。")
cur = p._p

p = caption(cur, parent, "表6-4 超分辨率重建质量对照表"); cur = p._p
sr = {r["scale"]: r for r in M["sr"]}
tbl = table_after(cur, ["放大倍数", "重建方法", "PSNR（dB）", "SSIM"], [
    ["2×", "双三次插值", f'{sr[2]["bicubic"]["psnr"]:.2f}', f'{sr[2]["bicubic"]["ssim"]:.4f}'],
    ["2×", "SRCNN", f'{sr[2]["srcnn"]["psnr"]:.2f}', f'{sr[2]["srcnn"]["ssim"]:.4f}'],
    ["3×", "双三次插值", f'{sr[3]["bicubic"]["psnr"]:.2f}', f'{sr[3]["bicubic"]["ssim"]:.4f}'],
    ["3×", "SRCNN", f'{sr[3]["srcnn"]["psnr"]:.2f}', f'{sr[3]["srcnn"]["ssim"]:.4f}'],
], [2.8, 4.0, 3.6, 3.6])
cur = tbl._tbl

p = body(cur, parent,
    "表6-4里的数据说明，SRCNN重建出来的峰值信噪比要比双三次插值高一点，2倍时高出0.5dB左右，3倍时高出0.6dB左右，"
    "在像素保真度上有一定优势；结构相似度方面两种方法咬得很紧，差距不到0.02，SRCNN甚至略微低了一点点。"
    "看得出在这个小规模训练集上训出来的SRCNN，细节增强带来的好处更多落在信噪比上，结构一致性并没有拉开明显差距。")
cur = p._p

p = body(cur, parent,
    "去噪效果用加噪前后的信噪比变化来刻画。做法是先往示例图的灰度图上分别叠加两种噪声，"
    "一种是标准差为20灰度级的高斯噪声，另一种是密度为5%的椒盐噪声，统计加噪后图像相对干净图像的信噪比；"
    "接着分别用均值滤波、高斯滤波、中值滤波和双边滤波去处理，再算出去噪后的信噪比。"
    "这里的信噪比定义成干净图像的方差除以残差（干净图像减去待测图像）的方差，再取10倍常用对数，单位是分贝。对照结果见表6-5。")
cur = p._p

p = caption(cur, parent, "表6-5 加噪前后信噪比对照表"); cur = p._p
S = M["snr"]
g, sp = S["高斯噪声σ=20"], S["椒盐噪声5%"]
tbl = table_after(cur, ["处理方式", "高斯噪声（σ=20）SNR（dB）", "椒盐噪声（5%）SNR（dB）"], [
    ["加噪后（未处理）", f'{g["加噪后"]:.2f}', f'{sp["加噪后"]:.2f}'],
    ["均值滤波", f'{g["均值滤波"]:.2f}', f'{sp["均值滤波"]:.2f}'],
    ["高斯滤波", f'{g["高斯滤波"]:.2f}', f'{sp["高斯滤波"]:.2f}'],
    ["中值滤波", f'{g["中值滤波"]:.2f}', f'{sp["中值滤波"]:.2f}'],
    ["双边滤波", f'{g["双边滤波"]:.2f}', f'{sp["双边滤波"]:.2f}'],
], [4.4, 5.0, 4.6])
cur = tbl._tbl

p = body(cur, parent,
    "从表6-5能读到几条信息：高斯噪声下，双边滤波的提升最明显，去噪后信噪比达到19.86dB，比加噪后的11.64dB高出8dB多，"
    "高斯滤波和中值滤波也有2dB上下的提升，均值滤波几乎没变化，甚至还略微掉了点，说明一味做均值平滑，噪声没压下去多少，图像的细节反而先被抹掉了一些；"
    "椒盐噪声下，中值滤波的表现最好，从5.58dB提到14.36dB，这正好印证了中值滤波对脉冲型噪声天生有优势这一点，"
    "而双边滤波在这类噪声下只有7.55dB，提升很有限，原因是孤立的脉冲点跟邻域灰度差太大，被它当成边缘给保留了下来。")
cur = p._p

p = body(cur, parent,
    "为了把关键参数对处理效果的影响讲清楚，测试环节还对几个代表性算法做了参数消融，也就是固定其他条件、只改一个参数，看结果指标怎么变。"
    "参与消融的有高斯滤波、双边滤波、伽马校正和Canny边缘检测这四种，扫描的分别是高斯核大小、双边颜色域标准差、伽马值和Canny低阈值，"
    "各条曲线的走向如图6-3所示。")
cur = p._p

p = picture(cur, parent, FIG); cur = p._p
p = caption(cur, parent, "图6-3 代表性算法关键参数取值-效果消融曲线"); cur = p._p

p = body(cur, parent,
    "图6-3给出几条规律：高斯滤波的核越大，去噪后的PSNR反而越低，核从3涨到15，PSNR从25.82dB掉到21.44dB，"
    "核太大了会把图像本身的细节也一并平滑掉；双边滤波的颜色域标准差存在一个比较合适的区间，在50附近PSNR最高，达到33.2dB，"
    "再往上调，保边效果还在，去噪能力却开始往下走；伽马值调大，输出图像的平均灰度单调往上抬，γ取0.3时平均灰度是149，取3.0时升到223.6，"
    "说明这个参数主要用来整体调亮度；Canny的低阈值调高，检出的边缘像素数逐步减少，阈值从30提到200，边缘点由38796个降到36423个，"
    "阈值越大，对边缘的判定就越严。")

# ---------- 同步前置清单 ----------
# 图清单(table0) / 表清单(table1) 加行
fig_list = doc.tables[0]
r = fig_list.add_row()
r.cells[0].text = "图6-3"; r.cells[1].text = "代表性算法关键参数取值-效果消融曲线"; r.cells[2].text = "36"

tab_list = doc.tables[1]
for num, name, pg in [("表6-4", "超分辨率重建质量对照表", "35"), ("表6-5", "加噪前后信噪比对照表", "36")]:
    r = tab_list.add_row()
    r.cells[0].text = num; r.cells[1].text = name; r.cells[2].text = pg

doc.save(PATH)
print("已保存:", PATH)
print("新表数:", len(doc.tables))
