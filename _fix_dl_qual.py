# -*- coding: utf-8 -*-
"""删除深度学习相关"效果得到验证"类主观定性表述
① 5.10 训练流水线有效性验证句
② 5.10 SRCNN 相对双三次插值的目视效果比较句
③ 6.2 表6-2 实际结果列 "与预期一致"
"""
from docx import Document

PATH = r'd:\1\bishe\毕设\7230267148_杨康康_基于opencv和streamlit的图像处理算法交互式演示系统 - 降重改写版.docx'

P1_KEY = "训练流水线是否有效，从这里就能得到验证"
P1_NEW = ("训练轮次不断增加的过程中，两个模型的损失逐步走低、准确率稳步爬升，具体变化曲线见图5-6。"
          "完整的训练流水线放在modules/training模块中，CIFAR10Net与SRCNN两个模型的训练都靠它支撑，"
          "数据加载与增强、模型构建、损失计算、反向传播与参数更新、验证评估、模型保存这些环节一个不少。"
          "细节安排上两者有分工：CIFAR-10数据集的训练图像先经随机裁剪和随机水平翻转完成数据增强，"
          "训练时拿交叉熵损失函数配Adam优化器；SRCNN那边则把CIFAR-10图像下采样，"
          "构造出低分辨率与高分辨率的训练样本对，损失函数改用均方误差。表5-8给出了深度学习算法的主要参数。")

P2_KEY = "SRCNN重建出的边缘与纹理细节更加锐利"
P2_NEW = ("把低分辨率图像重建为高分辨率图像，是图像超分辨率算法的目标，该算法建立在SRCNN网络之上。"
          "SRCNN用的是端到端的浅层卷积结构，低分辨率到高分辨率的非线性映射由它直接学习，"
          "算得上深度学习超分辨率的开创性工作之一。"
          "处理过程分两步，低分辨率输入先经双三次插值放大到目标尺寸，随后送入网络做细节增强，"
          "重建结果如图5-5(a)所示。SRCNN需要针对特定任务训练，为此系统附带了相应的训练脚本，"
          "训练一结束就可以在平台里做推理。")

RESULT_NEW = "输出尺寸与类型符合接口约定"


def set_text(p, text):
    if not p.runs:
        p.add_run(text)
        return
    p.runs[0].text = text
    for r in list(p.runs[1:]):
        r._element.getparent().remove(r._element)


doc = Document(PATH)

hit = []
for p in doc.paragraphs:
    if P1_KEY in p.text:
        set_text(p, P1_NEW)
        hit.append("①")
    elif P2_KEY in p.text:
        set_text(p, P2_NEW)
        hit.append("②")

rows_changed = 0
for tb in doc.tables:
    if not tb.rows:
        continue
    hdr = [c.text.strip() for c in tb.rows[0].cells]
    if hdr[:2] == ["测试项", "测试内容"] and hdr[-1] == "实际结果":
        for row in tb.rows[1:]:
            last = row.cells[-1]
            if last.text.strip() == "与预期一致":
                set_text(last.paragraphs[0], RESULT_NEW)
                rows_changed += 1

doc.save(PATH)
print(f"段落命中: {hit}, 表格行替换: {rows_changed}, 已保存")
