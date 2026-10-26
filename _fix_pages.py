# -*- coding: utf-8 -*-
"""用 Word COM 取题注显示页码(Information(1))，并把前置图/表清单页码校正为实测值"""
import sys
sys.stdout.reconfigure(encoding="utf-8")
from docx import Document
import win32com.client as wc

PATH = r'd:\1\bishe\毕设\7230267148_杨康康_基于opencv和streamlit的图像处理算法交互式演示系统 - 降重改写版.docx'

app = wc.Dispatch("Word.Application"); app.Visible = False
d = app.Documents.Open(PATH, ReadOnly=True)
pages = {}
try:
    for para in d.Paragraphs:
        if para.Range.Information(12):      # 跳过表格单元格
            continue
        t = para.Range.Text.replace("\r", "").replace("\x07", "").strip()
        for pre in ("图", "表"):
            if len(t) > 4 and t[0] == pre and "-" in t[:6]:
                key = t.split()[0].rstrip("．.")
                if key not in pages and len(key) >= 4:
                    pages[key] = para.Range.Information(1)
finally:
    d.Close(False); app.Quit()

print("实测显示页码:", pages)

doc = Document(PATH)
changed = []
for li, tname in ((0, "图"), (1, "表")):
    tl = doc.tables[li]
    for row in tl.rows[1:]:
        c = row.cells
        key = c[0].text.strip()
        if key not in pages:
            print(f"!! 未测到 {key}")
            continue
        newpg = str(pages[key])
        old = c[2].text.strip()
        if old != newpg:
            p = c[2].paragraphs[0]
            if p.runs:
                p.runs[0].text = newpg
                for r in p.runs[1:]:
                    r._element.getparent().remove(r._element)
            else:
                p.add_run(newpg)
            changed.append((tname, key, old, newpg))

doc.save(PATH)
print("已校正条目:")
for t in changed:
    print("   ", t)
print("共计校正", len(changed), "条")
