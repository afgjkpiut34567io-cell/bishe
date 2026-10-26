# -*- coding: utf-8 -*-
"""校验 6.4 追加内容与前置清单"""
import sys
sys.stdout.reconfigure(encoding="utf-8")
from docx import Document
PATH = r'd:\1\bishe\毕设\7230267148_杨康康_基于opencv和streamlit的图像处理算法交互式演示系统 - 降重改写版.docx'
doc = Document(PATH)
print(f"段落数={len(doc.paragraphs)} 表格数={len(doc.tables)}")

ps = doc.paragraphs
start = next(i for i, p in enumerate(ps) if "教学辅助的角度" in p.text)
end = next(i for i, p in enumerate(ps) if p.text.strip().startswith("本章小结") and i > start)
print(f"--- 6.4 区域 [{start}..{end}] ---")
for i in range(start, end + 1):
    t = ps[i].text
    tag = ""
    if "表6-4" in t or "表6-5" in t: tag = " <TAB>"
    if "图6-3" in t: tag += " <FIG>"
    if "PSNR" in t or "SSIM" in t or "SNR" in t or "消融" in t: tag += " <METRIC>"
    print(f"[{i}] ({ps[i].style.name}) {t[:80]}{tag}")

print("--- 新增表格 ---")
for ti in range(17, len(doc.tables)):
    tb = doc.tables[ti]
    print(f"TABLE[{ti}] rows={len(tb.rows)} cols={len(tb.columns)}")
    for r in tb.rows:
        print("   ", [c.text.strip() for c in r.cells])

print("--- 图清单末2行 ---")
for r in doc.tables[0].rows[-2:]:
    print("   ", [c.text.strip() for c in r.cells])
print("--- 表清单末3行 ---")
for r in doc.tables[1].rows[-3:]:
    print("   ", [c.text.strip() for c in r.cells])

# 图片是否嵌入
import zipfile
z = zipfile.ZipFile(PATH)
imgs = [n for n in z.namelist() if n.startswith("word/media/")]
print("--- word/media 图片数:", len(imgs))
print([n for n in imgs if "fig6" in n or n.endswith(".png")][-5:])
