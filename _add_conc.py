# -*- coding: utf-8 -*-
"""
诉求3：并发规模与部署边界
 1) 3.4 可行性分析：限定"浏览器即可访问"为演示场景 + 追加并发规模与部署边界段
 2) 6.3 性能测试：追加并发实测段 + 新增表6-4（并发实测结果表）
 3) 表6-4→表6-5、表6-5→表6-6 顺延（题注+正文引用）
 4) 7章创新点：部署优势限定为演示场景
 5) 前置表清单：插入表6-4、改名旧条目
"""
import sys, copy, shutil, json
sys.stdout.reconfigure(encoding="utf-8")
from docx import Document
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

PATH = r'd:\1\bishe\毕设\7230267148_杨康康_基于opencv和streamlit的图像处理算法交互式演示系统 - 降重改写版.docx'
BAK = r'd:\1\bishe\毕设\_bak_降重改写版_20261014_before_conc.docx'
M = json.load(open(r'd:\1\bishe\_conc_metrics.json', encoding='utf-8'))
shutil.copy2(PATH, BAK)
print("备份:", BAK)

doc = Document(PATH)


def find(sub):
    for p in doc.paragraphs:
        if sub in p.text:
            return p
    return None


def para_after(el, parent):
    p_el = OxmlElement('w:p')
    el.addnext(p_el)
    return Paragraph(p_el, parent)


def body(el, parent, text):
    p = para_after(el, parent)
    p.add_run(text)
    p.paragraph_format.first_line_indent = Pt(24)
    p.paragraph_format.line_spacing = 1.5
    return p


def caption(el, parent, text):
    p = para_after(el, parent)
    p.add_run(text)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.5
    return p


def table_after(el, header, rows, widths):
    tbl = doc.add_table(rows=len(rows) + 1, cols=len(header))
    tbl.style = doc.styles['Table Grid']
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    el.addnext(tbl._tbl)
    for j, h in enumerate(header):
        c = tbl.rows[0].cells[j]; c.text = ''
        cc = c.paragraphs[0]; cc.add_run(h); cc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = tbl.rows[i + 1].cells[j]; c.text = ''
            cc = c.paragraphs[0]; cc.add_run(str(v)); cc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in tbl.rows:
        for j, w in enumerate(widths):
            r.cells[j].width = Cm(w)
    return tbl


def clone_body(after_p, text):
    el = copy.deepcopy(after_p._p)
    after_p._p.addnext(el)
    np_ = Paragraph(el, after_p._parent)
    rpr = None
    for r in after_p.runs:
        f = r._element.find(qn('w:rPr'))
        if f is not None:
            rpr = copy.deepcopy(f); break
    for r in list(np_.runs):
        r._element.getparent().remove(r._element)
    run = np_.add_run(text)
    if rpr is not None:
        run._element.insert(0, rpr)
    return np_


def set_cell_text(cell, text):
    p = cell.paragraphs[0]
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r._element.getparent().remove(r._element)
    else:
        p.add_run(text)


# ============ 1) 3.4 ============
a34 = find("从技术可行性这方面来说")
assert a34 is not None, "未找到 3.4 锚点"
for r in a34.runs:
    if "用户只要把浏览器打开就可以访问" in r.text:
        r.text = r.text.replace("用户只要把浏览器打开就可以访问",
                                "在单机或者局域网演示场景里，用户只要把浏览器打开就可以访问")
        break
else:
    raise SystemExit("3.4 '浏览器即可访问' 定位失败")

T34 = ("再从并发规模和部署边界这个角度看一下，本系统定位在单用户演示或者小范围试用的层面，"
       "运行形态就是单用户本地运行；Streamlit这套框架靠的是一个服务进程把全部会话管起来，"
       "各个用户的会话状态分开存放，模型这类公共资源用cache_resource在进程里常驻一份，"
       "框架自己并不提供账号、权限或者多用户管理这些机制，所以系统适合一个人用，"
       "也适合课堂上几个人同时打开、轮着看效果，跟面向多用户并发的生产级服务不是一回事；"
       "服务在本机或者局域网里跑起来以后，同一网段里别的机器拿浏览器就能连上，"
       "“浏览器打开就能访问”这份便利，只在单机或者局域网演示这类小范围场景里才能成立，"
       "一旦并发访问的人数上去，界面刷新和算法计算就会互相抢占资源，响应时间也会明显被拉长；"
       "因此本文把系统的并发规模界定在单用户或者少量用户的演示场景，多用户并发访问并不算在它的能力范围之内。")
clone_body(a34, T34)

# ============ 2) 6.3 ============
b63 = find("会话状态数据始终保持着一致")
assert b63 is not None, "未找到 6.3 稳定性锚点"
parent = b63._parent

P1 = ("除了算法本体的耗时和长时间运行的稳定性之外，测试环节还补了一轮并发访问的实测，"
      "分成服务层和计算层两块来看。服务层测的是页面请求，同一时刻分别以1路、5路、10路和20路并发去访问系统的首页，"
      "把每次请求的耗时记下来；计算层测的是算法计算，同一时刻分别以1路、2路、4路和8路并发去调用高斯滤波，"
      "按默认参数处理800×600的图，把单次耗时和整体吞吐记下来。测试在一台16逻辑核心的CPU机器上进行，"
      "正式计时之前先做了预热，把首次页面编译、算法缓存这类一次性的开销排除掉，"
      "每一档都重复几轮之后取统计值，具体结果见表6-4。")

p = clone_body(b63, P1); cur = p._p
p = caption(cur, parent, "表6-4 并发实测结果表"); cur = p._p

svc = {s["concurrency"]: s for s in M["server"]}
comp = {c["concurrency"]: c for c in M["compute"]}
rows = []
for n in (1, 5, 10, 20):
    s = svc[n]
    rows.append([f"服务层（页面访问）", str(n), f'{s["avg_ms"]:.1f}', f'{s["max_ms"]:.1f}',
                 f'{s["ok"] / s["total"] * 100:.0f}%', "—"])
for n in (1, 2, 4, 8):
    c = comp[n]
    rows.append([f"计算层（高斯滤波）", str(n), f'{c["avg_ms"]:.2f}', f'{c["max_ms"]:.2f}',
                 "100%", f'{c["throughput"]:.1f}'])
tbl = table_after(cur, ["测试层次", "并发路数", "平均耗时（ms）", "最大耗时（ms）", "成功率", "吞吐（次/秒）"],
                  rows, [3.2, 1.8, 2.6, 2.6, 1.8, 2.8])
cur = tbl._tbl

P2 = ("从表6-4的数据能看到，服务层这边，并发路数从1路涨到20路，首页请求的平均耗时从8.2毫秒升到60.9毫秒，"
      "单次最大也没超过71.1毫秒，整个过程没有请求失败，说明在这个量级的并发下，页面分发这一环还撑得住；"
      "计算层这边，高斯滤波的单次耗时在1路时是0.71毫秒，4路时涨到2.05毫秒，8路时涨到3.13毫秒，"
      "单次延迟随着并发缓慢往上走，吞吐量则从1路的910次每秒提到8路的1651次每秒，"
      "也就是说多核确实被打进去了一部分，可并发越高单次等待越久，交互上的实时感受也会跟着打折。"
      "把这两块数据跟前文3.4小节讲过的部署边界合起来看，系统在单用户或者少量用户"
      "（画面里几个到十几个并发访问）的场景下能正常演示，并发规模再往上抬，响应时间就会明显拉长，"
      "所以它并不适合拿来当多用户并发的服务用。")
p = para_after(tbl._tbl, parent)
p.add_run(P2)
p.paragraph_format.first_line_indent = Pt(24)
p.paragraph_format.line_spacing = 1.5

# ============ 3) 表号顺延 表6-5→表6-6，再 表6-4→表6-5 ============
M65 = ["去噪效果用加噪前后的信噪比变化", "表6-5 加噪前后信噪比对照表", "从表6-5能读到"]
M64 = ["超分辨率这一块", "表6-4 超分辨率重建质量对照表", "表6-4里的数据说明"]
tag = {}
n_65 = n_64 = 0
for p in doc.paragraphs:
    t = p.text
    if "表6-5" in t and any(t.startswith(m) for m in M65):
        for r in p.runs:
            r.text = r.text.replace("表6-5", "表6-6"); n_65 += 1
    elif "表6-4" in t and any(t.startswith(m) for m in M64):
        tag[id(p)] = p
for p in tag.values():
    for r in p.runs:
        r.text = r.text.replace("表6-4", "表6-5"); n_64 += 1
print(f"重编号：表6-5->表6-6 {n_65} run | 表6-4->表6-5 {n_64} run")

# ============ 4) 7章创新点限定 ============
p321 = find("本系统的创新点体现在四个方面")
assert p321 is not None
for r in p321.runs:
    if r.text == "等优点":
        r.text = "等优点，不过这类好处主要落在单机或者小范围的演示场景里"
        break
else:
    raise SystemExit("创新点 '等优点' 定位失败")

# ============ 5) 前置表清单 ============
tl = doc.tables[1]
print("表清单行数(前):", len(tl.rows))
r13, r14 = tl.rows[13], tl.rows[14]
assert r13.cells[0].text.strip() == "表6-4", r13.cells[0].text
assert r14.cells[0].text.strip() == "表6-5", r14.cells[0].text
set_cell_text(r13.cells[0], "表6-5")
set_cell_text(r14.cells[0], "表6-6")

new = tl.add_row()
tl.rows[12]._tr.addnext(new._tr)
set_cell_text(new.cells[0], "表6-4")
set_cell_text(new.cells[1], "并发实测结果表")
set_cell_text(new.cells[2], "36")
print("表清单行数(后):", len(tl.rows))

doc.save(PATH)
print("已保存:", PATH)
for row in tl.rows:
    print("  ", [c.text.strip() for c in row.cells])
