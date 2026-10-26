# -*- coding: utf-8 -*-
"""5.10 追加预训练权重要说明段 + 表5-1 补一行"""
import sys, copy
sys.stdout.reconfigure(encoding="utf-8")
from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn

PATH = r'd:\1\bishe\毕设\7230267148_杨康康_基于opencv和streamlit的图像处理算法交互式演示系统 - 降重改写版.docx'
doc = Document(PATH)

anchor = None
for p in doc.paragraphs:
    if "这部分开销并没有计入后续的交互响应" in p.text:
        anchor = p; break
assert anchor is not None, "未找到 5.10 锚点"

P1 = ("图像分类和语义分割这两个算法，用的都是torchvision自带的预训练权重，来源是PyTorch官方托管站点，"
      "权重地址形如download.pytorch.org/models/再加文件名。具体到文件，MobileNetV2对应的是mobilenet_v2-7ebf99e0.pth，"
      "大小约13.6MB，换算成字节是14258573；语义分割用的DeepLabV3以MobileNetV3-Large作骨干，"
      "权重文件是deeplabv3_mobilenet_v3_large-fc3c493d.pth，大小约42.3MB，也就是44356159字节。"
      "这两份权重第一次被加载的时候，得先从磁盘读进来、再搭好网络结构，在纯CPU环境下实测，"
      "MobileNetV2的首次加载耗时约0.2秒，DeepLabV3约0.35秒；同一进程里再加载一次，因为权重已经在内存中，"
      "耗时分别降到0.16秒和0.34秒上下。文件名末尾那串十六进制字符是权重的哈希前缀，"
      "torchvision读文件时会拿它做完整性校验，因此文件名不能随便改。")

P2 = ("断网或者离线部署的情况下，加载是怎么走的，这里也说明一下。torchvision在加载权重时，"
      "会先去本地缓存目录里找同名文件，这个目录默认在用户目录下的.cache/torch/hub/checkpoints，"
      "本机没有设TORCH_HOME环境变量，走的就是默认路径。只要这些文件之前下载过一次，之后启动就不再需要联网，"
      "直接读本地缓存就行，断网也能正常加载；只有本地既没有缓存、又连不上网的时候，加载才会失败，"
      "这时候主程序的处理层会把异常接住，用错误提示把原因反馈给用户，同时回退到原图，整个系统并不会因此崩掉。"
      "另外还有两处细节：一是ImageNet的一千个类别名称也做了本地缓存，存在data目录下，只有本地文件缺失时才会尝试联网下载，"
      "下载失败就返回空，分类结果退化成类别编号，处理流程不会中断；二是深度学习依赖的torch和torchvision如果压根没装，"
      "算法注册表里的try/except会让这几个深度学习算法整体跳过，经典算法照常能用，"
      "这也是深度学习依赖被单独拆到requirements-dl.txt里的原因。")

def clone_body(after_p, text):
    el = copy.deepcopy(after_p._p)
    after_p._p.addnext(el)
    np_ = Paragraph(el, after_p._parent)
    rpr = None
    for r in after_p.runs:
        f = r._element.find(qn('w:rPr'))
        if f is not None:
            rpr = copy.deepcopy(f); break
    for r in list(np_.runs):
        r._element.getparent().remove(r._element)
    run = np_.add_run(text)
    if rpr is not None:
        run._element.insert(0, rpr)
    return np_

p1 = clone_body(anchor, P1)
p2 = clone_body(p1, P2)

# 表5-1 补一行
t51 = doc.tables[5]
print("表5-1 原行数:", len(t51.rows), "表头:", [c.text.strip() for c in t51.rows[0].cells])
r = t51.add_row()
r.cells[0].text = "深度学习框架"
r.cells[1].text = "PyTorch / torchvision（CPU 版）"

doc.save(PATH)
print("已保存。表5-1 新行数:", len(t51.rows))
for row in t51.rows:
    print("  ", [c.text.strip() for c in row.cells])
