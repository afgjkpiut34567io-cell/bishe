# -*- coding: utf-8 -*-
import sys
sys.stdout.reconfigure(encoding="utf-8")
from docx import Document

P = r'd:\1\bishe\毕设\7230267148_杨康康_基于opencv和streamlit的图像处理算法交互式演示系统 - 降重改写版.docx'
doc = Document(P)
KW = ["浏览器", "部署", "并发", "多用户", "单用户", "局域网", "分享", "访问", "表6-4", "表6-5", "表6-6"]
for i, p in enumerate(doc.paragraphs):
    t = p.text
    hit = [k for k in KW if k in t]
    if hit:
        print(f"[{i}] {p.style.name} {hit}")
        print("     ", t.strip()[:400])
print("===== 表格内命中 =====")
for ti, tb in enumerate(doc.tables):
    for ri, row in enumerate(tb.rows):
        for ci, c in enumerate(row.cells):
            for k in KW:
                if k in c.text:
                    print(f"table{ti} r{ri}c{ci} [{k}] {c.text.strip()[:80]}")
