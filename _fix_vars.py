# -*- coding: utf-8 -*-
"""[审查项#16] 变量注释表删去 η（学习率）、epochs（训练轮次）、batch_size（批大小）三行"""
import sys, shutil
sys.stdout.reconfigure(encoding="utf-8")
from docx import Document

PATH = r'd:\1\bishe\毕设\7230267148_杨康康_基于opencv和streamlit的图像处理算法交互式演示系统 - 降重改写版.docx'
BAK = r'd:\1\bishe\毕设\_bak_降重改写版_20261014_before_vars.docx'
shutil.copy2(PATH, BAK)
print("备份:", BAK)

doc = Document(PATH)

# 定位变量注释表（列头为 符号/名称/单位/物理意义）
tb = None
for t in doc.tables:
    hdr = [c.text.strip() for c in t.rows[0].cells]
    if hdr[:2] == ["符号", "名称"]:
        tb = t; break
assert tb is not None, "未找到变量注释表"
print("改前行数:", len(tb.rows))

DROP = {"η", "epochs", "batch_size"}
removed = []
for row in list(tb.rows):
    sym = row.cells[0].text.strip()
    if sym in DROP:
        removed.append(sym)
        row._tr.getparent().remove(row._tr)

doc.save(PATH)
print("已删除:", removed)
doc2 = Document(PATH)
for t in doc2.tables:
    if [c.text.strip() for c in t.rows[0].cells][:2] == ["符号", "名称"]:
        print("改后行数:", len(t.rows))
        for r in t.rows:
            print(" |", [c.text.strip() for c in r.cells])
