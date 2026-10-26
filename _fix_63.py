# -*- coding: utf-8 -*-
"""6.3 统一为"12种代表性算法"; 摘要按表6-3均值/中位数/样本数标注,使口径可复算"""
from docx import Document

PATH = r'd:\1\bishe\毕设\7230267148_杨康康_基于opencv和streamlit的图像处理算法交互式演示系统 - 降重改写版.docx'

EDITS = [
    # 6.3 表前引述段
    ("是部分代表性算法在800×600测试图像上的平均处理耗时",
     "是12种代表性算法在800×600测试图像上的平均处理耗时"),
    ("测试环节对12种算法逐一做了测量",
     "测试环节对这12种代表性算法逐一做了测量"),
    # 6.3 结果段
    ("绝大多数算法的平均处理耗时，都位于在1至4毫秒之间",
     "这12种代表性算法中，绝大多数算法的平均处理耗时都落在1至4毫秒之间"),
    # 6.3 分辨率对比段
    ("所测量的十二种算法的耗时在各档的分辨率下",
     "所测量的这12种代表性算法的耗时在各档的分辨率下"),
    # 中文摘要 —— 标注为表6-3均值,补中位数与样本数
    ("12种代表性算法的平均处理耗时约为3.7ms，中位数约为1.8ms，低于300ms这一算法核心计算耗时的指标要求，算是达标",
     "12种代表性算法在800×600测试图上的平均处理耗时约为3.7ms、中位数约为1.8ms（数据见表6-3，样本数为12），低于300ms这一算法核心计算耗时的指标要求，算是达标"),
    # 英文摘要 —— 同步
    ("Tests show that the twelve representative algorithms take about 3.7 ms on average (median about 1.8 ms) in algorithm core computation",
     "Tests show that the twelve representative algorithms take about 3.7 ms on average and about 1.8 ms at the median on the 800×600 test image (see Table 6-3, n = 12) in algorithm core computation"),
]


def set_text(p, text):
    if not p.runs:
        p.add_run(text)
        return
    p.runs[0].text = text
    for r in list(p.runs[1:]):
        r._element.getparent().remove(r._element)


doc = Document(PATH)
done, missed = [], list(EDITS)
for p in doc.paragraphs:
    t = p.text
    changed = False
    for pair in list(missed):
        old, new = pair
        if old in t:
            t = t.replace(old, new)
            done.append(old[:20])
            missed.remove(pair)
            changed = True
    if changed:
        set_text(p, t)

doc.save(PATH)
print(f"已应用 {len(done)} 处:")
for d in done:
    print("  -", d)
print(f"未命中 {len(missed)} 处:")
for old, _ in missed:
    print("  !", old[:30])
