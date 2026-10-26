# -*- coding: utf-8 -*-
"""合并pass5结果,修正映射,应用到文档"""
from docx import Document
import json, glob, re

SRC = r'd:\1\bishe\毕设\7230267148_杨康康_基于opencv和streamlit的图像处理算法交互式演示系统 - 降重改写版2.docx'
DST = r'd:\1\bishe\毕设\7230267148_杨康康_基于opencv和streamlit的图像处理算法交互式演示系统 - 降重改写版3.docx'

doc = Document(SRC)
paras = doc.paragraphs

# 报告(2)标红idx列表
report_red = [51, 82, 84, 91, 92, 93, 112, 117, 120, 121, 123, 125, 132, 138, 181, 182, 183, 185, 186, 209, 232, 260, 263, 277, 278, 296, 297, 298, 300, 301, 308, 309, 319, 323, 329, 331, 332, 343, 344, 345, 346, 348]

# 建立报告idx->文档idx映射(用文本前20字探针)
rep = Document(r'd:\download\免费_Word标红版_AIGC检测报告_[7230267148_杨康康_基于ope] (2).docx')
doc_texts = [p.text.strip() for p in doc.paragraphs]
rep_texts = [p.text.strip() for p in rep.paragraphs]

mapping = {}
for rid in report_red:
    rtext = rep_texts[rid]
    # 优先直接匹配
    if rid < len(doc_texts) and doc_texts[rid] and doc_texts[rid][:20] == rtext[:20]:
        mapping[str(rid)] = rid
        continue
    # 探针
    found = None
    for di, dt in enumerate(doc_texts):
        if dt and dt[:20] == rtext[:20]:
            found = di
            break
    mapping[str(rid)] = found if found is not None else rid

# 载入pass5输出
outs = {}
for f in sorted(glob.glob(r'd:\1\bishe\_rew\pass5_out*.json')):
    outs.update(json.load(open(f, encoding='utf-8')))

print(f"pass5输出段数: {len(outs)}, 映射覆盖: {len(mapping)}")

# 核验
FORBIDDEN = ['我觉得', '我认为', '我选了', '做的时候', '我在设计时', '我', '综上所述', '总而言之', '值得注意的是', '不难发现', '说白了']
for k, v in outs.items():
    if k not in mapping:
        print(f"!! 无映射: {k}")
        continue
    did = mapping[k]
    o = paras[did].text.strip()
    # 字数比
    ratio = len(v) / max(len(o), 1)
    if not (0.8 <= ratio <= 1.2):
        print(f"字数越界 [{k}->{did}] 原{len(o)} 新{len(v)} {ratio:.2f}")
    # 红线词
    for w in FORBIDDEN:
        if w in v:
            print(f"红线词 [{k}->{did}] {w}")
    # 引用图表
    for pat in [r'\[\d+\]', r'图\d+(?:-\d+)?', r'表\d+(?:-\d+)?']:
        if sorted(re.findall(pat, o)) != sorted(re.findall(pat, v)):
            print(f"标记不一致 [{k}->{did}] {pat}")

# 应用
applied = 0
for k, v in outs.items():
    did = mapping[k]
    p = paras[did]
    if not p.runs:
        print(f"跳过无run: {k}->{did}")
        continue
    p.runs[0].text = v
    for r in list(p.runs[1:]):
        r._element.getparent().remove(r._element)
    applied += 1

doc.save(DST)
print(f"已应用{applied}段, 保存: {DST}")

# 总字数
def total(path):
    d = Document(path)
    c = sum(len(re.sub(r'\s', '', p.text)) for p in d.paragraphs)
    t = sum(len(re.sub(r'\s', '', c2.text)) for tb in d.tables for row in tb.rows for c2 in row.cells)
    return c + t
print(f"源文件总字数 {total(SRC)}, 新文件总字数 {total(DST)}")
