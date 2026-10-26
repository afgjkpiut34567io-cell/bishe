# -*- coding: utf-8 -*-
"""补齐6.3表前引述段口径,并全文复检300ms相关表述"""
from docx import Document

PATH = r'd:\1\bishe\毕设\7230267148_杨康康_基于opencv和streamlit的图像处理算法交互式演示系统 - 降重改写版.docx'
OLD = "这组数据是出自实时性验证环节的：为了确认系统的实时性到底达没达标"
NEW = "这组数据是出自性能指标验证环节的：为了确认算法本体的计算耗时到底达没达标"


def set_text(p, text):
    if not p.runs:
        p.add_run(text)
        return
    p.runs[0].text = text
    for r in list(p.runs[1:]):
        r._element.getparent().remove(r._element)


doc = Document(PATH)
hit = False
for p in doc.paragraphs:
    if OLD in p.text:
        set_text(p, p.text.replace(OLD, NEW))
        hit = True
if hit:
    doc.save(PATH)

print("补齐:", hit)
print("=== 复检:仍含 实时性 / 300 的段落 ===")
for i, p in enumerate(doc.paragraphs):
    t = p.text
    if "实时性" in t or "300" in t:
        print(f"[{i}] {t}")
        print("-" * 70)
