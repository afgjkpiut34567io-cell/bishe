# -*- coding: utf-8 -*-
import sys
sys.stdout.reconfigure(encoding="utf-8")
from docx import Document

P = r'd:\1\bishe\毕设\7230267148_杨康康_基于opencv和streamlit的图像处理算法交互式演示系统 - 降重改写版.docx'
doc = Document(P)
KEYS = ["从技术可行性这方面来说", "会话状态数据始终保持着一致", "对照结果见表6-4",
        "表6-4 超分辨率", "表6-4里的数据", "对照结果见表6-5", "表6-5 加噪",
        "从表6-5能读到", "本系统的创新点体现在四个方面"]
for i, p in enumerate(doc.paragraphs):
    for k in KEYS:
        if k in p.text:
            print(f"[{i}] key={k} style={p.style.name} runs={len(p.runs)}")
            for j, r in enumerate(p.runs):
                print(f"    run{j} bold={r.bold} size={r.font.size} text={r.text[:60]!r}")
            print()
            break
print("==== 表清单 cell runs ====")
tl = doc.tables[1]
for ri in (12, 13, 14):
    print("row", ri, [c.text.strip() for c in tl.rows[ri].cells])
