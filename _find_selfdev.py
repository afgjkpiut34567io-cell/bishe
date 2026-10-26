# -*- coding: utf-8 -*-
import sys
sys.stdout.reconfigure(encoding="utf-8")
from docx import Document
P = r'd:\1\bishe\毕设\7230267148_杨康康_基于opencv和streamlit的图像处理算法交互式演示系统 - 降重改写版.docx'
doc = Document(P)
KW = ["自研", "自行实现", "自主实现", "自己实现", "独立实现", "四个", "四类", "均由", "都是自",
      "自己设计", "自发研制", "自主设计", "从零训练", "统一训练", "自行训练", "训练脚本", "训练流水线"]
for i, p in enumerate(doc.paragraphs):
    t = p.text
    hits = [k for k in KW if k in t]
    if hits:
        print(f"[{i}] {hits}")
        print("    ", t.strip()[:300])
for ti, tb in enumerate(doc.tables):
    for ri, row in enumerate(tb.rows):
        for c in row.cells:
            for k in KW:
                if k in c.text:
                    print(f"table{ti} r{ri}: [{k}] {c.text.strip()[:100]}")
                    break
