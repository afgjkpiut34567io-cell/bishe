# -*- coding: utf-8 -*-
"""打印前置 图清单(table0) 与 表清单(table1) 全部内容, 确认页码列是否填充"""
from docx import Document
PATH = r'd:\1\bishe\毕设\7230267148_杨康康_基于opencv和streamlit的图像处理算法交互式演示系统 - 降重改写版.docx'
doc = Document(PATH)
for ti in (0, 1):
    tb = doc.tables[ti]
    print(f"--- TABLE[{ti}] rows={len(tb.rows)} ---")
    for ri, row in enumerate(tb.rows):
        cells = [c.text.strip() for c in row.cells]
        print(f"  r{ri}: {cells}")
# 定位 6.4 区域锚点段落索引
print("--- 6.4 区域 ---")
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if t.startswith("教学辅助的角度") or t.startswith("这轮测试也让") or t.startswith("结果分析") or t.startswith("本章小结"):
        print(f"[{i}] {t[:40]}")
