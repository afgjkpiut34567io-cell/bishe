# -*- coding: utf-8 -*-
"""校正 表清单 中新增两行的页码（物理页-12）"""
import sys
sys.stdout.reconfigure(encoding="utf-8")
from docx import Document
PATH = r'd:\1\bishe\毕设\7230267148_杨康康_基于opencv和streamlit的图像处理算法交互式演示系统 - 降重改写版.docx'
doc = Document(PATH)
tl = doc.tables[1]
# 末两行为 表6-4 / 表6-5
tl.rows[-2].cells[2].text = "35"
tl.rows[-1].cells[2].text = "35"
doc.save(PATH)
for r in tl.rows[-3:]:
    print([c.text.strip() for c in r.cells])
for r in doc.tables[0].rows[-2:]:
    print([c.text.strip() for c in r.cells])
