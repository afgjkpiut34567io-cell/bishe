"""
毕业论文文档编辑脚本

对模板文档执行以下完善：
  1) 插入 10 张插图（图2-1 ~ 图6-2）及图题、正文引用；
  2) 补齐 5 张参数表（表5-3 ~ 表5-7）及表题、正文引用；
  3) 填充图清单、表清单；
  4) 重构并填充变量注释表；
  5) 更新摘要中的图/表数量声明；
  6) 清理孤立的空表。

用法：
    python build_thesis_docx.py
"""

import sys

sys.stdout.reconfigure(encoding="utf-8")

from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from docx.table import Table, _Cell

DOC_PATH = "毕设/（先用1）科文学院本科生毕业论文模板-中文模板（理工科).docx"
FIG_DIR = "figures"

# ==================== 基础工具 ====================

def find_para(doc, substr):
    """返回正文中第一个包含 substr 的段落（用于定位锚点）。"""
    for p in doc.paragraphs:
        if substr in p.text:
            return p
    raise RuntimeError(f"未找到锚点段落: {substr!r}")


def insert_paragraph_after(paragraph):
    """在指定段落之后插入一个空段落，返回新的 Paragraph 对象。"""
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    return Paragraph(new_p, paragraph._parent)


def set_centered(para, line_spacing=1.5):
    """设置段落居中、1.5 倍行距。"""
    pf = para.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf.line_spacing = line_spacing


def add_caption_after(anchor, text):
    """在 anchor 之后插入一条居中的题注段落（图题/表题）。"""
    p = insert_paragraph_after(anchor)
    run = p.add_run(text)
    set_centered(p)
    return p


def add_picture_after(anchor, img_path, width_cm=14.0):
    """在 anchor 之后插入一张居中的图片段落，返回该段落。"""
    p = insert_paragraph_after(anchor)
    run = p.add_run()
    run.add_picture(img_path, width=Cm(width_cm))
    set_centered(p)
    return p


def _set_cell_borders(cell):
    """为单元格补充单线边框（与占位表风格一致，备用）。"""
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tcPr.append(borders)


def make_table_after(anchor, header, rows, style="Table Grid"):
    """在 anchor 之后插入一张表格，返回 Table。"""
    tbl = doc.add_table(rows=len(rows) + 1, cols=len(header))
    tbl.style = doc.styles[style]
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    # 移动到 anchor 之后
    anchor._p.addnext(tbl._tbl)
    # 表头
    for j, h in enumerate(header):
        _set_text(tbl.rows[0].cells[j], h, bold=False)
    # 数据
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            _set_text(tbl.rows[i + 1].cells[j], val)
    _set_col_widths(tbl)
    return tbl


def _set_col_widths(tbl):
    """为表格设置合理列宽（三列参数表：窄/中/宽）。"""
    ncol = len(tbl.columns)
    widths = {2: [Cm(5), Cm(9)], 3: [Cm(3.2), Cm(4.6), Cm(6.2)],
              4: [Cm(2.6), Cm(3.2), Cm(2.6), Cm(5.6)]}
    ws = widths.get(ncol, [Cm(14.0 / ncol)] * ncol)
    for row in tbl.rows:
        for j, w in enumerate(ws):
            row.cells[j].width = w


def _set_text(cell, text, bold=False):
    """清空单元格并写入文本（默认字体，匹配正文）。"""
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold


def append_reference(para, text):
    """在段落末尾追加一句引用说明（继承正文格式）。"""
    run = para.add_run(text)
    return run


# ==================== 数据定义 ====================

FIGURES = [
    # (图号, 图名, 插图锚点子串, 引用段子串, 引用文本, 文件名)
    ("图2-1", "数字图像矩阵表示与二维卷积运算示意图",
     "卷积是空间域滤波的数学基础", "卷积是空间域滤波的数学基础",
     "数字图像矩阵表示与二维卷积运算的示意如图2-1所示。",
     "fig2_1_卷积示意.png"),
    ("图3-1", "系统总体用例图",
     "为更清晰地刻画系统与用户之间的交互", "为更清晰地刻画系统与用户之间的交互",
     "系统总体用例图如图3-1所示。",
     "fig3_1_系统用例图.png"),
    ("图4-1", "系统四层架构图",
     "各层之间的调用遵循单向依赖原则", "本系统采用模块化四层架构",
     "系统四层架构如图4-1所示。",
     "fig4_1_四层架构图.png"),
    ("图4-2", "系统数据流图",
     "会话状态在此过程中承担数据总线职能", "会话状态在此过程中承担数据总线职能",
     "系统数据流如图4-2所示。",
     "fig4_2_数据流图.png"),
    ("图4-3", "系统界面布局示意图",
     "这种布局设计的合理性在于", "系统界面采用Streamlit经典的",
     "系统界面布局如图4-3所示。",
     "fig4_3_界面布局图.png"),
    ("图5-1", "系统内置示例图像",
     "这种多元素合成图像的巧妙之处在于", "为方便用户在未准备素材的情况下",
     "系统内置示例图像如图5-1所示。",
     "fig5_1_示例图像.png"),
    ("图5-2", "滤波与图像增强算法处理效果对比",
     "def gamma_correction", "图像增强模块实现了直方图均衡化",
     "滤波与图像增强算法的处理效果对比如图5-2所示。",
     "fig5_2_滤波与增强对比.png"),
    ("图5-3", "边缘检测与形态学操作处理效果对比",
     "结构元素作为形态学运算的探针", "形态学操作以集合思想为基础",
     "边缘检测与形态学操作的处理效果对比如图5-3所示。",
     "fig5_3_边缘与形态学对比.png"),
    ("图6-1", "各算法平均处理耗时对比图",
     "测试结果显示，绝大多数算法的平均处理耗时", "测试结果显示，绝大多数算法的平均处理耗时",
     "各算法平均处理耗时对比如图6-1所示。",
     "fig6_1_耗时对比图.png"),
    ("图6-2", "处理耗时随图像分辨率变化曲线",
     "为进一步观察算法耗时与图像分辨率之间的关系", "为进一步观察算法耗时与图像分辨率之间的关系",
     "处理耗时随图像分辨率的变化关系如图6-2所示。",
     "fig6_2_耗时分辨率关系.png"),
]

PARAM_TABLES = [
    # (表号, 表名, 表头, 数据行, 插图锚点, 引用段, 引用文本)
    ("表5-3", "图像增强算法参数表", ["算法", "参数", "含义"], [
        ["直方图均衡化", "—", "无参数，自动计算累积分布函数"],
        ["CLAHE", "clip_limit / tile_grid_size", "对比度限制阈值 / 分块尺寸"],
        ["对比度亮度调整", "alpha / beta", "对比度增益 / 亮度偏移"],
        ["伽马校正", "gamma", "伽马值 γ，控制灰度非线性映射"],
        ["锐化", "strength", "锐化强度"],
        ["反锐化掩膜", "sigma / strength", "高斯模糊标准差 / 增强强度"],
    ], "def gamma_correction", "图像增强模块实现了直方图均衡化",
     "图像增强算法的主要参数如表5-3所示。"),
    ("表5-4", "边缘检测算法参数表", ["算法", "参数", "含义"], [
        ["Canny", "threshold1 / threshold2", "低阈值 / 高阈值"],
        ["Sobel", "ksize / dx / dy", "核尺寸 / x方向导数阶 / y方向导数阶"],
        ["Laplacian", "ksize", "核尺寸"],
        ["Scharr", "dx / dy", "x方向导数 / y方向导数"],
        ["Prewitt", "—", "无参数，固定 3×3 差分核"],
    ], "Sobel与Prewitt算子的差异集中于卷积核", "边缘检测模块实现了Canny",
     "边缘检测算法的主要参数如表5-4所示。"),
    ("表5-5", "形态学操作参数表", ["算法", "参数", "含义"], [
        ["腐蚀", "ksize / iterations", "结构元素尺寸 / 迭代次数"],
        ["膨胀", "ksize / iterations", "结构元素尺寸 / 迭代次数"],
        ["开运算", "ksize", "结构元素尺寸"],
        ["闭运算", "ksize", "结构元素尺寸"],
        ["形态学梯度", "ksize", "结构元素尺寸"],
        ["顶帽", "ksize", "结构元素尺寸"],
        ["黑帽", "ksize", "结构元素尺寸"],
    ], "结构元素作为形态学运算的探针", "形态学操作以集合思想为基础",
     "形态学操作的主要参数如表5-5所示。"),
    ("表5-6", "阈值分割算法参数表", ["算法", "参数", "含义"], [
        ["全局阈值", "thresh / maxval", "分割阈值 / 最大值"],
        ["自适应阈值", "block_size / C", "邻域尺寸 / 偏移常数"],
        ["Otsu二值化", "—", "无参数，自动确定最优阈值"],
    ], "在OpenCV中，Otsu方法通过类型参数", "阈值分割模块实现了全局阈值",
     "阈值分割算法的主要参数如表5-6所示。"),
    ("表5-7", "图像变换算法参数表", ["算法", "参数", "含义"], [
        ["缩放", "scale", "缩放比例"],
        ["旋转", "angle / scale", "旋转角度 / 缩放比例"],
        ["仿射变换", "angle / scale / tx / ty", "旋转角度 / 缩放 / x平移 / y平移"],
        ["傅里叶变换", "—", "无参数，输出频谱图"],
    ], "傅里叶变换的实现流程分为以下几个步骤", "图像变换模块涵盖几何变换",
     "图像变换算法的主要参数如表5-7所示。"),
]

VARIABLES = [
    ["f", "图像矩阵", "灰度级", "表示数字图像的二维矩阵，元素为像素灰度值"],
    ["w", "卷积核", "无量纲", "空间域滤波的权重矩阵，又称滤波模板"],
    ["g", "输出图像", "灰度级", "卷积运算后的结果矩阵"],
    ["γ", "伽马系数", "无量纲", "伽马校正的幂指数，控制灰度非线性映射"],
    ["σ", "标准差", "像素", "高斯分布标准差，决定高斯核的离散程度"],
    ["t", "阈值", "灰度级", "像素分类的判定阈值"],
    ["C", "偏移常数", "灰度级", "自适应阈值中邻域均值减去的常数"],
    ["d", "滤波直径", "像素", "双边滤波的邻域直径"],
    ["ksize", "核尺寸", "像素", "卷积核的空间支撑范围（边长）"],
    ["block_size", "邻域尺寸", "像素", "自适应阈值计算的邻域窗口大小"],
    ["σ_color", "颜色域标准差", "灰度级", "双边滤波颜色相似度权重参数"],
    ["σ_space", "空间域标准差", "像素", "双边滤波空间邻近度权重参数"],
]


# ==================== 主流程 ====================

def main():
    global doc
    doc = Document(DOC_PATH)

    # ---- 1. 插入插图 + 图题（图题在图下方） ----
    print("插入插图...")
    for (figno, figname, ins_anchor, ref_anchor, ref_text, fname) in FIGURES:
        a = find_para(doc, ins_anchor)
        pic = add_picture_after(a, f"{FIG_DIR}/{fname}")
        add_caption_after(pic, f"{figno}  {figname}")
        print(f"  {figno} {figname}")

    print("插入参数表...")
    for (tbno, tbname, header, rows, ins_anchor, ref_anchor, ref_text) in PARAM_TABLES:
        a = find_para(doc, ins_anchor)
        cap = add_caption_after(a, f"{tbno}  {tbname}")
        make_table_after(cap, header, rows)
        print(f"  {tbno} {tbname}")

    # ---- 2. 正文引用（追加到引用段落）----
    print("添加正文引用...")
    for (figno, figname, ins_anchor, ref_anchor, ref_text, fname) in FIGURES:
        p = find_para(doc, ref_anchor)
        append_reference(p, ref_text)
    for (tbno, tbname, header, rows, ins_anchor, ref_anchor, ref_text) in PARAM_TABLES:
        p = find_para(doc, ref_anchor)
        append_reference(p, ref_text)

    # ---- 3. 填充图清单 / 表清单 ----
    print("填充图清单 / 表清单...")
    fill_list(doc.tables[0], FIGURES, kind="图")
    fill_list(doc.tables[1], build_table_list(), kind="表")

    # ---- 4. 变量注释表 ----
    print("重构变量注释表...")
    build_variable_table(doc.tables[2])

    # ---- 5. 删除孤立空表 ----
    print("清理孤立空表...")
    remove_stray_table(doc)

    # ---- 6. 更新摘要 ----
    print("更新摘要数量声明...")
    update_abstract(doc)

    # 保存到新文件（避免与 Word 中打开的源文档锁冲突）
    out = DOC_PATH.replace(".docx", "_完善.docx")
    doc.save(out)
    print(f"\n已保存: {out}")


def build_table_list():
    """按出现顺序构建表清单条目（图号+名称）。"""
    return [
        ("表4-1", "算法分类统计表"),
        ("表5-1", "开发环境配置表"),
        ("表5-2", "滤波算法参数表"),
        ("表5-3", "图像增强算法参数表"),
        ("表5-4", "边缘检测算法参数表"),
        ("表5-5", "形态学操作参数表"),
        ("表5-6", "阈值分割算法参数表"),
        ("表5-7", "图像变换算法参数表"),
        ("表6-1", "功能测试用例表"),
        ("表6-2", "部分算法平均处理耗时（800×600图像）"),
    ]


def fill_list(table, entries, kind):
    """填充图清单/表清单。entries: [(编号, 名称), ...] 或 FIGURES 结构。"""
    # 归一化 entries 为 (编号, 名称)
    if kind == "图":
        norm = [(e[0], e[1]) for e in entries]
    else:
        norm = entries
    n_data = len(norm)
    # 调整数据行数（保留表头）
    cur_rows = len(table.rows)
    need_rows = n_data + 1
    if need_rows > cur_rows:
        for _ in range(need_rows - cur_rows):
            table.add_row()
    elif need_rows < cur_rows:
        for _ in range(cur_rows - need_rows):
            table._tbl.remove(table.rows[-1]._tr)
    # 填充
    for i, (num, name) in enumerate(norm):
        row = table.rows[i + 1]
        _set_text(row.cells[0], num)
        _set_text(row.cells[1], name)
        _set_text(row.cells[2], "")  # 页码待定


def build_variable_table(table):
    """把 2 列占位表扩展为 4 列变量注释表并填充。"""
    # 现有列数
    ncols = len(table.columns)
    for _ in range(4 - ncols):
        table.add_column(Cm(3.0))
    header = ["符号", "名称", "单位", "物理意义"]
    for j, h in enumerate(header):
        _set_text(table.rows[0].cells[j], h, bold=False)
    for i, var in enumerate(VARIABLES):
        for j, val in enumerate(var):
            _set_text(table.rows[i + 1].cells[j], val)
    _set_col_widths(table)


def remove_stray_table(doc):
    """删除孤立的空表（1 行 × 2 列、无内容）。"""
    for tbl in list(doc.tables):
        if len(tbl.rows) == 1 and len(tbl.columns) == 2:
            if all(c.text.strip() == "" for c in tbl.rows[0].cells):
                tbl._tbl.getparent().remove(tbl._tbl)
                print("  已删除孤立空表")
                return


def update_abstract(doc):
    """更新中英文摘要的图/表数量声明。"""
    pairs = [
        ("该论文有图15幅，表26个，参考文献160篇",
         "该论文有图10幅，表10个，参考文献160篇"),
        ("This paper contains 15 figures, 26 tables, and 160 references",
         "This paper contains 10 figures, 10 tables, and 160 references"),
    ]
    for old, new in pairs:
        for p in doc.paragraphs:
            if old in p.text:
                # 重写段落（保留首 run 格式）
                fmt_run = p.runs[0] if p.runs else None
                for r in list(p.runs):
                    r._element.getparent().remove(r._element)
                run = p.add_run(new)
                if fmt_run is not None:
                    run.font.name = fmt_run.font.name
                    run.font.size = fmt_run.font.size
                    run.font.bold = fmt_run.font.bold
                print(f"  已更新: {new[:20]}...")
                break


if __name__ == "__main__":
    main()
