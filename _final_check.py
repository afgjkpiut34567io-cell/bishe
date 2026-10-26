# -*- coding: utf-8 -*-
import sys
sys.stdout.reconfigure(encoding="utf-8")
from docx import Document
P = r'd:\1\bishe\毕设\7230267148_杨康康_基于opencv和streamlit的图像处理算法交互式演示系统 - 降重改写版.docx'
doc = Document(P)
print("段落总数", len(doc.paragraphs), "| 表格总数", len(doc.tables))
print("\n--- 图清单 ---")
for row in doc.tables[0].rows:
    print("  ", [c.text.strip() for c in row.cells])
print("\n--- 表清单 ---")
for row in doc.tables[1].rows:
    print("  ", [c.text.strip() for c in row.cells])
print("\n--- 新增并发表 ---")
for ti, tb in enumerate(doc.tables):
    head = [c.text.strip() for c in tb.rows[0].cells]
    if head and head[0] == "测试层次":
        print("table index", ti, "rows", len(tb.rows))
        for row in tb.rows:
            print("  ", [c.text.strip() for c in row.cells])
print("\n--- 6.3 尾部/6.4 段落顺序 ---")
ps = doc.paragraphs
flag = False
for i, p in enumerate(ps):
    t = p.text.strip()
    if t.startswith("表6-3给出的"):
        flag = True
    if flag and t:
        print(f"[{i}] {t[:80]}")
    if t.startswith("本章围绕系统把功能测试"):
        break
