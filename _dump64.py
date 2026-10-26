# -*- coding: utf-8 -*-
"""导出目标文档 6.2~6.5 区域段落文本 + 表格清单, 用于定位补充指标的位置"""
from docx import Document
PATH = r'd:\1\bishe\毕设\7230267148_杨康康_基于opencv和streamlit的图像处理算法交互式演示系统 - 降重改写版.docx'
doc = Document(PATH)

print(f"段落数: {len(doc.paragraphs)}, 表格数: {len(doc.tables)}")

# 找 6.2 / 6.3 / 6.4 / 6.5 / 第7章 标题位置
anchors = ["6.1", "6.2", "6.3", "6.4", "6.5", "第7章", "7.1"]
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    for a in anchors:
        if t.startswith(a) or t.startswith(a.replace(".", ".")):
            print(f"HEAD [{i}] ({p.style.name}) {t[:60]}")
            break

print("=" * 60)
start = None
for i, p in enumerate(doc.paragraphs):
    if p.text.strip().startswith("6.2"):
        start = i
        break
if start is None:
    start = 260
for i in range(start, min(start + 90, len(doc.paragraphs))):
    p = doc.paragraphs[i]
    txt = p.text
    mark = ""
    if "图" in txt and "图6" in txt:
        mark = " <FIG>"
    if "表6" in txt:
        mark += " <TAB6>"
    print(f"[{i}] ({p.style.name}) {txt}{mark}")

print("=" * 60)
for ti, tb in enumerate(doc.tables):
    hdr = " | ".join(c.text.strip()[:14] for c in tb.rows[0].cells)
    print(f"TABLE[{ti}] rows={len(tb.rows)} cols={len(tb.columns)} hdr= {hdr}")
