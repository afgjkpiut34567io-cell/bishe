# -*- coding: utf-8 -*-
import sys
sys.stdout.reconfigure(encoding="utf-8")
from docx import Document
P = r'd:\1\bishe\毕设\7230267148_杨康康_基于opencv和streamlit的图像处理算法交互式演示系统 - 降重改写版.docx'
doc = Document(P)
print("段落", len(doc.paragraphs), "表格", len(doc.tables))
ps = doc.paragraphs
flag = False
for i, p in enumerate(ps):
    t = p.text.strip()
    if t.startswith("从表6-4的数据能看到"):
        flag = True
    if flag and t:
        print(f"[{i}] {t[:95]}")
    if t.startswith("本章围绕系统把功能测试"):
        break
print("\n--- 新内存表 ---")
for ti, tb in enumerate(doc.tables):
    h = [c.text.strip() for c in tb.rows[0].cells]
    if h and h[0] == "模型":
        print("index", ti)
        for r in tb.rows:
            print("  ", [c.text.strip() for c in r.cells])
print("\n--- 图6-3/6-4、表6-5/6-6/6-7 引用 ---")
for i, p in enumerate(ps):
    t = p.text
    for k in ("图6-3", "图6-4", "表6-5", "表6-6", "表6-7"):
        if k in t:
            print(f"[{i}] {k}: {t.strip()[:70]}")
            break
