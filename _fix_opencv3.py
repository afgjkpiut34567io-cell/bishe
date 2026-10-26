# -*- coding: utf-8 -*-
"""[审查项#17 同源] 4.1 架构设计 / 5.11 本章小结 / 7.1 工作总结 三处 34种整体归 OpenCV 的口径统一"""
import sys, shutil
sys.stdout.reconfigure(encoding="utf-8")
from docx import Document

PATH = r'd:\1\bishe\毕设\7230267148_杨康康_基于opencv和streamlit的图像处理算法交互式演示系统 - 降重改写版.docx'
BAK = r'd:\1\bishe\毕设\_bak_降重改写版_20261014_before_opencv3.docx'
shutil.copy2(PATH, BAK)
print("备份:", BAK)

doc = Document(PATH)

# ---------- [150] 4.1 架构设计 ----------
p = doc.paragraphs[150]
assert p.runs[4].text == "完成了七大类共" and p.runs[5].text == "34"
p.runs[4].text = "完成了六大类"
p.runs[5].text = "30"
p.runs[6].text = "种经典算法，另外四类深度学习算法则交给PyTorch来完成，系统的核心计算由它承担；可视化层基于"
print("\n[150]:", p.text)

# ---------- [268] 5.11 本章小结 ----------
p = doc.paragraphs[268]
assert p.runs[1].text == "34"
assert p.runs[0].text.endswith("七大类")
p.runs[0].text = p.runs[0].text[: -len("七大类")] + "六大类"
p.runs[1].text = "30"
p.runs[2].text = "种经典算法的处理功能，是由"
p.runs[4].text = "这个核心来完成的，四类深度学习算法则靠PyTorch来支撑，交互式界面是由"
print("\n[268]:", p.text)

# ---------- [335] 7.1 工作总结 ----------
p = doc.paragraphs[335]
assert p.runs[3].text == "34" and p.runs[2].text == "把"
p.runs[2].text = "把六大类"
p.runs[3].text = "30"
p.runs[4].text = "种经典算法、基于PyTorch把四类深度学习算法一个一个做了出来，再拿"
print("\n[335]:", p.text)

doc.save(PATH)
print("\n已保存。")
